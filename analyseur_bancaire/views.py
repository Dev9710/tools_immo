# views.py - Version refactorisée avec templates
from pathlib import Path
from urllib.parse import urlencode

from django.shortcuts import render, redirect
from django.http import HttpResponse, JsonResponse
from django.template.loader import render_to_string
from django.conf import settings
from django.urls import reverse
from datetime import datetime, date, timedelta
from dateutil.relativedelta import relativedelta
from decimal import Decimal, InvalidOperation
import math
import tempfile
import os
import re
import unicodedata
import dataclasses
import json
import subprocess
import time

from . import biens_gino

# RÈGLE : toujours les taux les plus récents possible (consigne utilisateur du
# 30/09/2026). Mettre à jour TAUX_DATE avec la table ; l'app affiche la date et
# alerte au-delà de BAREME_VALIDITE_JOURS.
#
# Taux nominaux hors assurance, profil « moyen » = profil « bon » de Meilleurtaux,
# baromètre du 05/10/2026 : 3,76 / 3,85 / 4,00 % sur 15 / 20 / 25 ans (excellent
# 3,30 / 3,40 / 3,50 ; très bon 3,55 / 3,68 / 3,75). Hausse continue, le 25 ans
# atteint 4 % (01/10 : 3,71 / 3,80 / 3,93 ; usure 5,40 % au 01/10 pour 20 ans et +).
# 7 et 10 ans : même écart au 15 ans qu'auparavant (−0,07 ; −0,02). Écarts régionaux
# conservés (IDF −0,08 ; Provence −0,03 ; Rhône-Alpes −0,06 par rapport au national).
TAUX_DATE = date(2026, 10, 5)
TAUX_SOURCE = "Meilleurtaux"
BAREME_VALIDITE_JOURS = 7    # Meilleurtaux publie son baromètre chaque semaine
TAUX_ACTUELS = {
    'regions': {
        'ile_de_france': {'7': 3.61, '10': 3.66, '15': 3.68, '20': 3.77, '25': 3.92},
        'provence': {'7': 3.66, '10': 3.71, '15': 3.73, '20': 3.82, '25': 3.97},
        'rhone_alpes': {'7': 3.63, '10': 3.68, '15': 3.70, '20': 3.79, '25': 3.94},
        'autre': {'7': 3.69, '10': 3.74, '15': 3.76, '20': 3.85, '25': 4.00}
    },
    'profils': {
        'excellent': -0.30,    # CDI, >10% apport, épargne
        'bon': -0.15,          # CDI, 10% apport
        'moyen': 0.00,         # CDD, apport minimal
        'risque': 0.25         # Profil difficile
    },
    # Assurance emprunteur en délégation, non-fumeur (barèmes 2026).
    'assurance': {
        'moins_30': 0.10,
        '30_45': 0.20,
        '45_plus': 0.40
    }
}

# Frais de notaire 2026. Toute l'Île-de-France a relevé les droits de mutation à
# 5 % au 01/01/2026 (Val-de-Marne : 6,32 % de DMTO) → ~8 % dans l'ancien. Les
# primo-accédants sont exonérés de la hausse → ~7,5 %.
def taux_pour_duree(duree, region='ile_de_france'):
    """Taux nominal « moyen » du barème pour une durée quelconque.

    Le barème ne donne que 7/10/15/20/25 ans : entre deux paliers on interpole,
    au-delà on garde le palier extrême. Plus c'est long, plus c'est cher.
    """
    table = sorted((int(k), v) for k, v in TAUX_ACTUELS['regions'][region].items())
    if duree <= table[0][0]:
        return table[0][1]
    for (d1, t1), (d2, t2) in zip(table, table[1:]):
        if duree <= d2:
            return round(t1 + (t2 - t1) * (duree - d1) / (d2 - d1), 2)
    return table[-1][1]


def bareme_durees(region='ile_de_france'):
    """{durée: taux} de 5 à 30 ans, pour ajuster le taux dès qu'on change la durée."""
    return {d: taux_pour_duree(d, region) for d in range(5, 31)}


def bareme_info(aujourd_hui=None):
    """Date et source du barème, et s'il est périmé (affichés sur les pages)."""
    age = ((aujourd_hui or date.today()) - TAUX_DATE).days
    return {'date': TAUX_DATE.strftime('%d/%m/%Y'), 'source': TAUX_SOURCE,
            'perime': age > BAREME_VALIDITE_JOURS, 'age_jours': age,
            'taux_20': f"{TAUX_ACTUELS['regions']['autre']['20']:.2f}".replace('.', ','),
            # Valeurs par défaut des champs de saisie (point décimal, pour <input type="number">)
            'champ_taux_20': f"{TAUX_ACTUELS['regions']['autre']['20']:.2f}",
            'champ_assurance': f"{TAUX_ACTUELS['assurance']['30_45']:.2f}"}


FRAIS_NOTAIRE = {
    'ancien': 0.08,
    'ancien_primo': 0.075,
    'neuf': 0.03
}


class SimulateurPretImmobilier:
    """Simulateur de prêt immobilier professionnel"""

    def __init__(self):
        self.taux_endettement_max = 35.0  # HCSF 2025
        self.duree_max = 27  # années

    def calculer_capacite_emprunt(self, revenus_nets, charges_mensuelles, duree,
                                  region='autre', profil='moyen', age=35):
        """Calcule la capacité d'emprunt selon les critères bancaires"""

        # Mensualité maximum : la règle HCSF plafonne le TOTAL des charges
        # (existantes + futur crédit) à 35 % des revenus nets. La mensualité
        # du crédit ne peut donc pas dépasser 35 % des revenus MOINS les
        # charges déjà supportées.
        mensualite_max = max(
            0.0, revenus_nets * (self.taux_endettement_max / 100) - charges_mensuelles)

        # Taux selon région et profil
        taux_base = TAUX_ACTUELS['regions'][region].get(str(duree), 3.5)
        ajustement_profil = TAUX_ACTUELS['profils'][profil]
        taux_nominal = taux_base + ajustement_profil

        # Taux assurance selon âge
        if age < 30:
            taux_assurance = TAUX_ACTUELS['assurance']['moins_30']
        elif age < 45:
            taux_assurance = TAUX_ACTUELS['assurance']['30_45']
        else:
            taux_assurance = TAUX_ACTUELS['assurance']['45_plus']

        # TAEG = taux nominal + assurance
        taeg = taux_nominal + taux_assurance

        # Calcul capacité d'emprunt (formule actuarielle)
        nb_mensualites = duree * 12
        if taeg == 0:
            capacite = mensualite_max * nb_mensualites
        else:
            taux_mensuel = taeg / 100 / 12
            capacite = mensualite_max * \
                (1 - (1 + taux_mensuel) ** -nb_mensualites) / taux_mensuel

        # Coût total et intérêts
        cout_total = mensualite_max * nb_mensualites
        cout_interets = cout_total - capacite

        # Taux d'endettement projeté = (charges existantes + mensualité du
        # crédit) / revenus. Plafonné à 35 % par le calcul ci-dessus.
        if revenus_nets > 0:
            taux_endettement = round(
                (charges_mensuelles + mensualite_max) / revenus_nets * 100, 2)
        else:
            taux_endettement = 0.0

        return {
            'capacite_emprunt': round(capacite, 2),
            'mensualite_max': round(mensualite_max, 2),
            'taux_nominal': round(taux_nominal, 2),
            'taux_assurance': round(taux_assurance, 2),
            'taeg': round(taeg, 2),
            'cout_total': round(cout_total, 2),
            'cout_interets': round(cout_interets, 2),
            'taux_endettement': taux_endettement
        }

    def calculer_mensualites(self, montant_emprunt, duree, taux_nominal, taux_assurance):
        """Calcule les mensualités pour un montant donné"""

        taeg = taux_nominal + taux_assurance
        taux_mensuel = taeg / 100 / 12
        nb_mensualites = duree * 12

        if taeg == 0:
            mensualite = montant_emprunt / nb_mensualites
        else:
            mensualite = montant_emprunt * taux_mensuel / \
                (1 - (1 + taux_mensuel) ** -nb_mensualites)

        cout_total = mensualite * nb_mensualites
        cout_interets = cout_total - montant_emprunt

        return {
            'mensualite': round(mensualite, 2),
            'cout_total': round(cout_total, 2),
            'cout_interets': round(cout_interets, 2),
            'taeg': round(taeg, 2)
        }

    def calculer_frais_notaire(self, prix_bien, type_bien='ancien'):
        """Calcule les frais de notaire selon le type de bien"""
        taux = FRAIS_NOTAIRE[type_bien]
        frais = prix_bien * taux
        return round(frais, 2)

    def generer_tableau_amortissement(self, montant, duree, taux_nominal, taux_assurance, nb_lignes=12):
        """Génère un tableau d'amortissement simplifié"""

        taeg = taux_nominal + taux_assurance
        taux_mensuel = taeg / 100 / 12
        nb_mensualites = duree * 12

        if taeg == 0:
            mensualite = montant / nb_mensualites
        else:
            mensualite = montant * taux_mensuel / \
                (1 - (1 + taux_mensuel) ** -nb_mensualites)

        tableau = []
        capital_restant = montant

        for i in range(min(nb_lignes, int(nb_mensualites))):
            if taeg == 0:
                interets = 0
                capital = mensualite
            else:
                interets = capital_restant * taux_mensuel
                capital = mensualite - interets

            capital_restant = max(0, capital_restant - capital)

            tableau.append({
                'mois': i + 1,
                'mensualite': round(mensualite, 2),
                'capital': round(capital, 2),
                'interets': round(interets, 2),
                'capital_restant': round(capital_restant, 2)
            })

        return tableau

    def calculer_reste_a_vivre(self, revenus_nets, charges_totales, nb_adultes=2, nb_enfants=0):
        """Calcule et valide le reste à vivre"""

        # Seuils minimums par personne (données 2025)
        seuil_adulte = 400  # €/mois
        seuil_enfant = 300  # €/mois

        minimum_requis = (nb_adultes * seuil_adulte) + \
            (nb_enfants * seuil_enfant)
        reste_a_vivre = revenus_nets - charges_totales

        alerte = None
        if reste_a_vivre < minimum_requis:
            deficit = minimum_requis - reste_a_vivre
            alerte = {
                'type': 'danger',
                'message': f"Reste à vivre insuffisant : {reste_a_vivre}€ < {minimum_requis}€ requis",
                'conseil': f"Réduisez vos charges de {deficit}€ ou augmentez vos revenus"
            }
        elif reste_a_vivre < minimum_requis * 1.2:
            alerte = {
                'type': 'warning',
                'message': f"Reste à vivre limite : {reste_a_vivre}€",
                'conseil': "Prévoyez une marge de sécurité supplémentaire"
            }

        return {
            'reste_a_vivre': reste_a_vivre,
            'minimum_requis': minimum_requis,
            'alerte': alerte,
            'statut': 'OK' if not alerte else alerte['type']
        }


class BanquePostaleParserSimple:
    """Parser Banque Postale optimisé et nettoyé"""

    def __init__(self):
        # Classification basée uniquement sur des mots-clés réalistes
        self.keywords = {
            'achat': [
                'ACHAT CB', 'PAIEMENT CB', 'CARTE', 'TPE', 'RETRAIT',
                'PAYPAL', 'UBER', 'INTERSPORT', 'IMMOBILIERE', 'CB '
            ],
            'virement': [
                'VIREMENT', 'VIR ', 'SEPA', 'REMBOURSEMENT',
                'INSTANTANE', 'EMISSION', 'RECEPTION', 'TRANSFERT'
            ],
            'depot': [
                'DEPOT', 'VERSEMENT', 'REMISE', 'CREDIT'
            ],
            'debit': [
                'PRELEVEMENT', 'PRLV', 'ECHEANCE', 'ABONNEMENT',
                'COTISATION', 'FRAIS', 'TRESOR', 'ORANGE', 'EDF'
            ],
            'cheque': [
                'CHEQUE', 'CHQ'
            ]
        }

    def extract_text_from_pdf(self, pdf_path):
        """Extrait le texte de toutes les pages du PDF"""
        try:
            import pdfplumber
            text = ""
            with pdfplumber.open(pdf_path) as pdf:
                for page_num, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text += f"\n--- PAGE {page_num + 1} ---\n{page_text}\n"
            return text
        except ImportError:
            # Fallback PyPDF2
            import PyPDF2
            text = ""
            with open(pdf_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page_num, page in enumerate(pdf_reader.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text += f"\n--- PAGE {page_num + 1} ---\n{page_text}\n"
            return text

    def parse_montant(self, montant_str):
        """Parse les montants français avec espaces (ex: 10 000,00)"""
        try:
            montant_clean = montant_str.replace(' ', '').replace(',', '.')
            return Decimal(montant_clean)
        except (InvalidOperation, ValueError, AttributeError):
            return Decimal('0')

    def categorize_operation(self, description):
        """Catégorise une opération selon les mots-clés"""
        desc_upper = description.upper()

        for keyword in self.keywords['cheque']:
            if keyword in desc_upper:
                return 'cheque'

        for keyword in self.keywords['debit']:
            if keyword in desc_upper:
                return 'debit'

        for keyword in self.keywords['achat']:
            if keyword in desc_upper:
                return 'achat'

        for keyword in self.keywords['virement']:
            if keyword in desc_upper:
                return 'virement'

        for keyword in self.keywords['depot']:
            if keyword in desc_upper:
                return 'depot'

        return 'autre'

    # Restes du type d'opération, sans valeur informative comme libellé.
    _FRAGMENT_RE = re.compile(
        r'(?:INSTANTANE|PERMANENT|IMMEDIAT)?\s*(?:A|DE|POUR|D)?', re.IGNORECASE)

    def extract_libelle(self, description):
        """Extrait le libellé principal optimisé pour Banque Postale"""
        libelle = description.strip()

        # Supprimer références inutiles
        libelle = re.sub(r'REF(?:ERENCE)?\s*:\s*\d+', '', libelle)
        libelle = re.sub(r'\b\d{10,}\b', '', libelle)  # Numéros longs
        libelle = re.sub(r'COMPTE\s+FR\d+', '', libelle)  # IBAN

        # Cas CB avec astérisque (*UBER B)
        match = re.search(
            r'(?:ACHAT|PAIEMENT)\s+CB\s+(.+?\*?.+?)(?:\s+\d{2}\.\d{2}|\s+CARTE|$)',
            libelle.upper()
        )
        if match:
            return match.group(1).replace('*', ' ').strip()

        # Cas CB classique
        match = re.search(
            r'(?:ACHAT|PAIEMENT)\s+CB\s+(.+?)(?:\s+\d{2}\.\d{2}|\s+CARTE|$)',
            libelle.upper()
        )
        if match:
            return match.group(1).strip()

        # Cas virements. Le relevé ne porte pas toujours le bénéficiaire sur la
        # ligne de l'opération : la capture peut alors ne ramener qu'un
        # fragment du type d'opération (« INSTANTANEA », « POUR »), moins
        # parlant que la description entière. On ne garde donc l'extraction que
        # si elle apporte vraiment un nom.
        if 'VIREMENT' in libelle.upper():
            patterns = [
                r'VIREMENT\s+(?:INSTANTANE\s+)?(?:A|POUR)\s+(.+?)(?:\s+COMPTE|\s+DEFAULT|$)',
                r'VIREMENT\s+(.+?)(?:\s+COMPTE|\s+REFERENCE|$)',
            ]
            for pattern in patterns:
                match = re.search(pattern, libelle.upper())
                if match:
                    candidat = match.group(1).strip()
                    if not self._FRAGMENT_RE.fullmatch(candidat):
                        return candidat
                    break

        # Cas prélèvements
        if 'PRELEVEMENT DE' in libelle.upper():
            match = re.search(
                r'PRELEVEMENT DE\s+(.+?)(?:\s+\d+|$)', libelle.upper())
            if match:
                return match.group(1).strip()

        # Cas chèques
        if 'CHEQUE N°' in libelle.upper():
            return libelle

        # Nettoyage final
        libelle = re.sub(r'\s+', ' ', libelle).strip()
        return libelle[:60]

    # Montant au format d'un mot isolé (« 17,00 », « 3360,00 », « 1 234,56 »).
    AMOUNT_RE = re.compile(r'^\d{1,3}(?:[ .]?\d{3})*,\d{2}$')
    _OPERATIONS_MARKER = re.compile(
        r'(?:Vos\s+op[eé]rations|OP[EÉ]RATIONS)', re.IGNORECASE)

    # « Relevé édité le 12 septembre 2025 » : le seul repère de date fiable du
    # document. pdfplumber colle souvent les mots, d'où les \s* partout.
    _EDITION_RE = re.compile(
        r'dit[ée]?\s*le\s*(\d{1,2})\s*'
        r'(janvier|f[ée]vrier|mars|avril|mai|juin|juillet|ao[uû]t|'
        r'septembre|octobre|novembre|d[ée]cembre)\s*(20\d{2})',
        re.IGNORECASE)
    _MOIS_NUM = {
        'janvier': 1, 'fevrier': 2, 'février': 2, 'mars': 3, 'avril': 4,
        'mai': 5, 'juin': 6, 'juillet': 7, 'aout': 8, 'août': 8,
        'septembre': 9, 'octobre': 10, 'novembre': 11,
        'decembre': 12, 'décembre': 12,
    }

    def extract_edition(self, text):
        """Date d'édition du relevé (« Relevé édité le 12 septembre 2025 »),
        ou None.

        Ne PAS se rabattre sur la première date jj/mm/aaaa du document : les
        mentions légales en contiennent (« PEL ouvert jusqu'au 31/12/2017 »),
        ce qui datait les relevés récents de plusieurs années.
        """
        m = self._EDITION_RE.search(text)
        if not m:
            return None
        mois = self._MOIS_NUM.get(m.group(2).lower())
        if not mois:
            return None
        try:
            return date(int(m.group(3)), mois, int(m.group(1)))
        except ValueError:
            return None

    # « du 12/02/2026 au 11/03/2026 » : quand le relevé porte sa période en
    # toutes lettres, c'est le repère le plus sûr.
    _PERIODE_RE = re.compile(
        r'du\s*(\d{2}/\d{2}/20\d{2})\s*au\s*(\d{2}/\d{2}/20\d{2})',
        re.IGNORECASE)

    def extract_periode(self, text):
        """(début, fin) de la période couverte par le relevé, ou None.

        Un relevé ne suit pas le mois civil : édité le 12 septembre, il court
        du 12 août au 11 septembre. C'est cette fenêtre — déjà longue d'un
        mois — qui sert d'unité de mesure, plutôt qu'un découpage en mois
        civils qui obligerait à fournir deux PDF pour un seul mois complet.
        """
        # Une mention légale peut contenir « du … au … » : on n'accepte le
        # motif que si l'intervalle ressemble à un mois de relevé.
        for brut_debut, brut_fin in self._PERIODE_RE.findall(text[:2000]):
            try:
                debut = datetime.strptime(brut_debut, '%d/%m/%Y').date()
                fin = datetime.strptime(brut_fin, '%d/%m/%Y').date()
            except ValueError:
                continue
            if 25 <= (fin - debut).days + 1 <= 35:
                return debut, fin

        edition = self.extract_edition(text)
        if not edition:
            return None
        return edition - relativedelta(months=1), edition - timedelta(days=1)

    def build_year_resolver(self, text):
        """Renvoie une fonction mois → année.

        Un relevé est à cheval sur deux mois (édité le 12/09, il couvre le
        12/08 au 11/09) et peut donc franchir un 31 décembre. Les opérations
        dont le mois dépasse celui de l'édition appartiennent à l'année
        précédente — sinon un relevé de janvier daterait décembre de l'année
        suivante.
        """
        edition = self.extract_edition(text)
        if edition:
            mois_ref, annee_ref = edition.month, edition.year
            return lambda mois: annee_ref if mois <= mois_ref else annee_ref - 1

        # Repli : « Relevé … 20aa » dans l'en-tête, sinon l'année courante.
        m = re.search(r'Relev[eé].*?(20\d{2})', text)
        annee = int(m.group(1)) if m else datetime.now().year
        return lambda mois: annee

    # « 8 157,80 » arrive de pdfplumber en deux mots : « 8 » puis « 157,80 ».
    # Ces deux motifs servent à les recoller.
    _MILLIERS_RE = re.compile(r'^\d{1,3}$')
    _CENTAINES_RE = re.compile(r'^\d{3}(?:[  .]?\d{3})*,\d{2}$')
    # Écart horizontal maximal, en points, entre deux fragments d'un même
    # nombre. Au-delà, ce sont deux valeurs distinctes.
    ECART_MAX_FRAGMENTS = 6

    # Abscisse de la frontière Débit/Crédit sur un relevé Banque Postale, quand
    # l'en-tête n'a pas pu être localisé.
    SEUIL_PAR_DEFAUT = 503.0
    # Les colonnes de montants occupent la moitié droite de la page : un mot
    # « crédit » situé à gauche appartient à une description d'opération
    # (« CREDIT CARTE BANCAIRE »), pas à l'en-tête du tableau.
    X_MIN_COLONNES = 300

    # « Total des opérations   8 157,80   10 839,32 » : le relevé porte lui-même
    # ses totaux débit et crédit. C'est le seul contrôle indépendant possible.
    _TOTAUX_RE = re.compile(r'(?i)total\s*des\s*op')

    def extract_totaux(self, lines):
        """(total débit, total crédit) imprimés sur le relevé, ou None.

        Sert à vérifier que le parsing n'a rien perdu. Sans ce garde-fou, un
        montant mal découpé passait inaperçu : les totaux paraissaient
        plausibles et le dossier partait à la banque avec de faux chiffres.
        """
        for ligne in lines:
            if self._TOTAUX_RE.search(ligne['text']) and len(ligne['amounts']) >= 2:
                montants = sorted(ligne['amounts'], key=lambda a: a[1])[:2]
                debit = self.parse_montant(montants[0][0])
                credit = self.parse_montant(montants[1][0])
                if debit or credit:
                    return debit, credit
        return None

    def _recoller_milliers(self, mots):
        """Fusionne les fragments d'un même montant.

        pdfplumber découpe « 8 157,80 » en « 8 » et « 157,80 » : sans
        recollage, seul « 157,80 » est reconnu comme montant et l'opération
        perd ses milliers. Le total d'un relevé s'en trouvait amputé de
        plusieurs milliers d'euros, sans aucun signe d'erreur.

        `mots` est trié par abscisse. On ne fusionne que si le fragment de
        gauche est un groupe de 1 à 3 chiffres, celui de droite un groupe de
        centaines complet, et les deux visuellement collés.
        """
        fusionnes = []
        for mot in mots:
            if fusionnes:
                gauche = fusionnes[-1]
                if (self._MILLIERS_RE.match(gauche['text'])
                        and self._CENTAINES_RE.match(mot['text'])
                        and mot['x0'] - gauche['x1'] < self.ECART_MAX_FRAGMENTS):
                    fusionnes[-1] = {**gauche,
                                     'text': gauche['text'] + ' ' + mot['text'],
                                     'x1': mot['x1']}
                    continue
            fusionnes.append(mot)
        return fusionnes

    def _bornes_entete(self, mots_de_la_ligne):
        """(x1 de « Débit », x1 de « Crédit ») si cette ligne est l'en-tête du
        tableau des opérations, sinon (None, None).

        Les deux mots doivent figurer sur la MÊME ligne et dans la zone des
        colonnes de montants. Sans ces deux conditions, le premier « CREDIT »
        rencontré dans un libellé d'opération devenait la frontière : le seuil
        tombait vers x=97 au lieu de x=503, et la totalité des montants était
        classée en crédit — un relevé entier disparaissait alors de l'analyse.
        """
        debit = credit = None
        for w in mots_de_la_ligne:
            if w['x1'] < self.X_MIN_COLONNES:
                continue
            t = w['text'].lower().strip(':')
            if debit is None and t.startswith(('débit', 'debit')):
                debit = w['x1']
            elif credit is None and t.startswith(('crédit', 'credit')):
                credit = w['x1']
        return (debit, credit) if debit and credit else (None, None)

    def extract_word_lines(self, pdf_path):
        """Reconstruit les lignes du relevé à partir des mots *positionnés* et
        détecte la frontière horizontale entre les colonnes Débit et Crédit.

        Chaque ligne renvoyée est un dict {'text', 'amounts'} où `amounts` liste
        les (valeur, x1) des montants de la ligne. Le sens d'une opération se lit
        ainsi à la position du montant, et non plus par mots-clés — ce qui évite
        de compter un crédit comme une charge. Renvoie (None, None) si pdfplumber
        est absent : on retombe alors sur le texte brut, sans notion de sens."""
        try:
            import pdfplumber
        except ImportError:
            return None, None

        lines = []
        debit_right = credit_right = None
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                words = page.extract_words()
                # Regrouper les mots en lignes visuelles (tolérance verticale).
                clusters = []
                for w in sorted(words, key=lambda w: w['top']):
                    if clusters and abs(w['top'] - clusters[-1]['top']) <= 3:
                        clusters[-1]['words'].append(w)
                    else:
                        clusters.append({'top': w['top'], 'words': [w]})

                for cl in clusters:
                    toks = self._recoller_milliers(
                        sorted(cl['words'], key=lambda w: w['x0']))
                    if debit_right is None or credit_right is None:
                        d, c = self._bornes_entete(toks)
                        if d and c:
                            debit_right, credit_right = d, c
                    lines.append({
                        'text': ' '.join(w['text'] for w in toks),
                        'amounts': [(w['text'], w['x1'])
                                    for w in toks if self.AMOUNT_RE.match(w['text'])],
                    })

        # Frontière Débit/Crédit : milieu des deux bords droits d'en-tête, sinon
        # la valeur de repli observée sur les relevés Banque Postale.
        if debit_right and credit_right:
            seuil = (debit_right + credit_right) / 2
        else:
            seuil = self.SEUIL_PAR_DEFAUT
        return lines, seuil

    _CIVILITE = re.compile(r'^(MR|MME|MLLE|M\.?|MONSIEUR|MADAME)\s+', re.IGNORECASE)

    def extract_titulaires(self, lignes):
        """Noms des titulaires lus dans l'en-tête (« MR DUPONT OU MME DUPONT MARTIN »).

        La 1re ligne à civilité est celle du titulaire ; celle du conseiller vient
        après. Sert à reconnaître les virements vers ses propres comptes sans
        jamais écrire un nom dans le code.
        """
        for ligne in lignes[:60]:
            ligne = ligne.strip()
            if not self._CIVILITE.match(ligne):
                continue
            noms = []
            for part in re.split(r'\s+(?:OU|ET)\s+', ligne.upper()):
                for mot in self._CIVILITE.sub('', part).split():
                    if len(mot) >= 4 and mot.isalpha() and mot not in noms:
                        noms.append(mot)
            return noms
        return []

    def _clean_description(self, description):
        """Nettoie une description : retire les dates parasites, recolle les
        libellés agglutinés et normalise les espaces."""
        description = re.sub(
            r'\d{2}[./]\d{2}[./](?:20)?\d{2}', '', description)
        description = description.replace("ACHATCB", "ACHAT CB ")
        # La Banque Postale exporte souvent les mots collés
        # ("VIREMENTINSTANTANEA") : on ré-espace les termes d'en-tête connus,
        # sinon le libellé affiché est illisible.
        description = re.sub(
            r'\b(VIREMENT|PRELEVEMENT|PAIEMENT|ACHAT|RETRAIT|CHEQUE|'
            r'INSTANTANE|REMBOURSEMENT|COTISATION)(?=[A-Z])',
            r'\1 ', description)
        description = re.sub(r'IMMOBILIERE(\d+)', r'IMMOBILIERE \1', description)
        return re.sub(r'\s+', ' ', description).strip()

    def parse_operations_lines(self, lines, seuil, annee_de):
        """Parse les lignes positionnées ; le sens (débit/crédit) vient de la
        colonne dans laquelle tombe le montant. `annee_de` est la fonction
        mois → année construite par build_year_resolver."""
        operations = []
        i = 0
        while i < len(lines):
            text = lines[i]['text'].strip()

            if not text or any(w in text for w in
                               ['Ancien solde', 'Nouveau solde', 'Solde', 'Débit', 'Crédit']):
                i += 1
                continue

            date_match = re.match(r'^(\d{2}/\d{2})\s+(.+)', text)
            if not date_match:
                i += 1
                continue

            jour_mois = date_match.group(1)
            try:
                mois_op = int(jour_mois.split('/')[1])
                date_operation = datetime.strptime(
                    f"{jour_mois}/{annee_de(mois_op)}", '%d/%m/%Y').date()
            except ValueError:
                i += 1
                continue

            description = date_match.group(2)
            amounts = lines[i]['amounts']

            # Montant absent de la ligne de date : chercher sur les suivantes,
            # jusqu'à la prochaine opération, en agrégeant la description.
            j = i
            while not amounts and j + 1 < len(lines):
                nxt = lines[j + 1]
                if re.match(r'^\d{2}/\d{2}\s+', nxt['text'].strip()):
                    break
                if nxt['amounts']:
                    amounts = nxt['amounts']
                    break
                if not any(s in nxt['text'] for s in
                           ['CARTE NUMERO', 'REF :', 'IDENT :', 'MANDAT :']):
                    description += ' ' + nxt['text']
                j += 1

            if not amounts:
                i += 1
                continue

            # Le bénéficiaire d'un virement est sur la ligne SUIVANTE (« DUPONT
            # JEAN Livret A ») : sans elle, impossible de distinguer un
            # virement vers son épargne d'une vraie dépense.
            complement = []
            if j == i:
                for nxt in lines[i + 1:i + 3]:
                    t = nxt['text'].strip()
                    if not t or nxt['amounts'] or re.match(r'^\d{2}/\d{2}\s+', t):
                        break
                    if not t.upper().startswith('REFERENCE'):
                        complement.append(t)

            # Une ligne d'opération porte un seul montant ; s'il y en a plusieurs
            # on prend le plus à droite (colonne des montants).
            value_str, x1 = max(amounts, key=lambda a: a[1])
            montant = self.parse_montant(value_str)
            sens = 'credit' if x1 > seuil else 'debit'

            description = self._clean_description(
                description.replace(value_str, ' '))

            if montant > Decimal('0'):
                operations.append({
                    'date': date_operation,
                    'description': description,
                    'montant': montant,
                    'sens': sens,
                    'type': self.categorize_operation(description),
                    'libelle': self.extract_libelle(description),
                    'complement': re.sub(r'\d{6,}', '', ' '.join(complement)).strip(),
                })
            i += 1

        return operations

    def parse_operations_section(self, text_section, annee_de=None):
        """Repli texte (sans position) : on ne connaît pas le sens débit/crédit,
        il reste à None. Utilisé seulement si pdfplumber est indisponible."""
        operations = []
        lines = text_section.split('\n')
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line or any(word in line for word in
                               ['Date', 'Operation', 'Debit', 'Credit',
                                'Ancien solde', 'Nouveau solde']):
                i += 1
                continue

            date_match = re.match(r'^(\d{2}/\d{2})\s+(.+)', line)
            if not date_match:
                i += 1
                continue

            jour_mois = date_match.group(1)
            try:
                mois_op = int(jour_mois.split('/')[1])
                annee = annee_de(mois_op) if annee_de else datetime.now().year
                date_operation = datetime.strptime(
                    f"{jour_mois}/{annee}", '%d/%m/%Y').date()
            except ValueError:
                i += 1
                continue

            description = self._clean_description(date_match.group(2))
            montant = Decimal('0')
            montant_matches = list(re.finditer(
                r'(\d{1,3}(?:\s\d{3})*,\d{2})', description))
            if montant_matches:
                last_match = montant_matches[-1]
                montant = self.parse_montant(last_match.group(1))
                description = (description[:last_match.start()]
                               + description[last_match.end():]).strip()
            else:
                j = i + 1
                while j < min(i + 4, len(lines)):
                    next_line = lines[j].strip()
                    if re.match(r'^\d{2}/\d{2}\s+', next_line):
                        break
                    if re.match(r'^\d{1,3}(?:\s\d{3})*,\d{2}$', next_line):
                        montant = self.parse_montant(next_line)
                        break
                    if next_line and not any(skip in next_line for skip in
                                             ['CARTE NUMERO', 'REF :', 'IDENT :', 'MANDAT :']):
                        description += ' ' + next_line
                    j += 1

            if montant > Decimal('0'):
                operations.append({
                    'date': date_operation,
                    'description': description,
                    'montant': montant,
                    'sens': None,
                    'type': self.categorize_operation(description),
                    'libelle': self.extract_libelle(description),
                })
            i += 1

        return operations

    def parse_pdf(self, pdf_path, types_selectionnes=None):
        """Parse complet du PDF Banque Postale (géométrie des colonnes en
        priorité, texte brut en repli)."""
        try:
            lines, seuil = self.extract_word_lines(pdf_path)

            if lines is not None:
                full_text = '\n'.join(l['text'] for l in lines)
                if len(full_text) < 50:
                    return {'success': False, 'error': 'PDF vide ou illisible'}
                annee_de = self.build_year_resolver(full_text)
                periode = self.extract_periode(full_text)
                totaux = self.extract_totaux(lines)

                # Début de la section : la ligne d'en-tête « Débit … Crédit »
                # (présente sur tout relevé) est la plus fiable ; sinon la
                # mention « opérations ». Si aucune n'est trouvée, on parcourt
                # tout et on se fie au repérage date + montant-en-colonne.
                header_idx = next(
                    (idx for idx, l in enumerate(lines)
                     if re.search(r'(?i)d[eé]bit', l['text'])
                     and re.search(r'(?i)cr[eé]dit', l['text'])), None)
                if header_idx is None:
                    header_idx = next(
                        (idx for idx, l in enumerate(lines)
                         if self._OPERATIONS_MARKER.search(l['text'])), None)

                start = 0 if header_idx is None else header_idx + 1
                all_operations = self.parse_operations_lines(
                    lines[start:], seuil, annee_de)

                if not all_operations and header_idx is None:
                    return {'success': False,
                            'error': 'Section "Vos opérations" non trouvée'}
            else:
                text = self.extract_text_from_pdf(pdf_path)
                if not text or len(text) < 50:
                    return {'success': False, 'error': 'PDF vide ou illisible'}
                if not self._OPERATIONS_MARKER.search(text):
                    return {'success': False, 'error': 'Section "Vos opérations" non trouvée'}
                annee_de = self.build_year_resolver(text)
                periode = self.extract_periode(text)
                totaux = None  # sans positions, la ligne de totaux est illisible
                sections = self._OPERATIONS_MARKER.split(text)
                all_operations = []
                for section in sections[1:]:
                    all_operations.extend(
                        self.parse_operations_section(section, annee_de))

            # Grouper par type ; en charges fixes on ne veut que les sorties :
            # un crédit (sens détecté par la colonne) est exclu même si un
            # mot-clé le rangeait par erreur dans « debit »/« achat ».
            # types_selectionnes=None : on garde tout, « autre » compris. C'est le
            # mode utilisé par le total des dépenses, où aucune sortie ne doit
            # être perdue à cause d'un libellé non reconnu.
            tous_types = list(self.keywords.keys()) + ['autre']
            if types_selectionnes is None:
                types_selectionnes = tous_types

            operations_par_type = {k: [] for k in tous_types}
            for op in all_operations:
                if op['type'] in types_selectionnes:
                    operations_par_type[op['type']].append({
                        'date': op['date'],
                        'libelle': op['libelle'],
                        'description': op['description'],
                        'complement': op.get('complement', ''),
                        'montant': op['montant'],
                        'sens': op.get('sens'),
                    })

            return {
                'success': True,
                'titulaires': (self.extract_titulaires([l['text'] for l in lines])
                               if lines is not None else []),
                'operations_par_type': operations_par_type,
                'periode': periode,
                'totaux_releve': totaux,
                'types_traites': [t for t in types_selectionnes if operations_par_type[t]]
            }

        except Exception as e:
            return {'success': False, 'error': f'Erreur parsing: {str(e)}'}


def create_excel_multi_onglets(operations_par_type, types_traites):
    """Crée un Excel avec un onglet par type d'opération"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return create_csv_multi_files(operations_par_type, types_traites)

    workbook = Workbook()
    workbook.remove(workbook.active)  # Supprimer feuille par défaut

    type_names = {
        'achat': 'Achats CB',
        'virement': 'Virements',
        'depot': 'Depots & Versements',
        'debit': 'Prelevements & Debits',
        'cheque': 'Cheques'
    }

    type_colors = {
        'achat': 'FF6B6B',
        'virement': '4ECDC4',
        'depot': '45B7D1',
        'debit': 'FFA07A',
        'cheque': 'DDA0DD'
    }

    for type_op in types_traites:
        if not operations_par_type[type_op]:
            continue

        # Créer onglet
        ws = workbook.create_sheet(type_names[type_op])
        ws.sheet_properties.tabColor = type_colors.get(type_op, '808080')

        # En-tête principal
        ws['A1'] = f"ANALYSE {type_names[type_op].upper()}"
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:D1')

        # Informations
        ws['A2'] = f"Nombre d'operations: {len(operations_par_type[type_op])}"
        ws['A2'].font = Font(size=12)

        ws['A3'] = f"Date d'analyse: {datetime.now().strftime('%d/%m/%Y a %H:%M')}"
        ws['A3'].font = Font(size=10, italic=True)

        # Headers tableau avec style
        headers = ['Date', 'Libelle', 'Description complete', 'Montant (€)']
        header_fill = PatternFill(
            start_color=type_colors[type_op], end_color=type_colors[type_op], fill_type="solid")
        border = Border(left=Side(style='thin'), right=Side(style='thin'),
                        top=Side(style='thin'), bottom=Side(style='thin'))

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        # Données avec alternance de couleurs
        row = 6
        total = Decimal('0')

        # Trier les opérations par date
        operations_triees = sorted(
            operations_par_type[type_op], key=lambda x: x['date'])

        for i, operation in enumerate(operations_triees):
            # Couleur alternée
            if i % 2 == 0:
                fill = PatternFill(start_color='F8F9FA',
                                   end_color='F8F9FA', fill_type="solid")
            else:
                fill = PatternFill(start_color='FFFFFF',
                                   end_color='FFFFFF', fill_type="solid")

            # Date
            cell = ws.cell(row=row, column=1,
                           value=operation['date'].strftime('%d/%m/%Y'))
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

            # Libellé
            cell = ws.cell(row=row, column=2, value=operation['libelle'])
            cell.fill = fill
            cell.border = border
            cell.font = Font(bold=True)

            # Description
            cell = ws.cell(row=row, column=3, value=operation['description'])
            cell.fill = fill
            cell.border = border

            # Montant (valeur absolue pour l'affichage)
            montant_val = float(abs(operation['montant']))
            cell = ws.cell(row=row, column=4, value=montant_val)
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = '#,##0.00'

            # Couleur selon le sens (colonne Débit/Crédit du relevé)
            if operation.get('sens') == 'credit':
                cell.font = Font(color="28A745")  # Vert pour les crédits
            else:
                cell.font = Font(color="DC3545")  # Rouge pour les débits

            total += abs(operation['montant'])  # Somme en valeur absolue
            row += 1

        # Ligne totale avec style
        row += 1
        ws.merge_cells(f'A{row}:C{row}')
        total_cell = ws.cell(row=row, column=1, value="TOTAL")
        total_cell.font = Font(bold=True, size=14)
        total_cell.alignment = Alignment(horizontal="center")
        total_cell.fill = PatternFill(
            start_color="FFD966", end_color="FFD966", fill_type="solid")
        total_cell.border = border

        # Montant total
        total_montant_cell = ws.cell(row=row, column=4, value=float(total))
        total_montant_cell.font = Font(bold=True, size=14)
        total_montant_cell.fill = PatternFill(
            start_color="FFD966", end_color="FFD966", fill_type="solid")
        total_montant_cell.border = border
        total_montant_cell.alignment = Alignment(horizontal="right")
        total_montant_cell.number_format = '#,##0.00 "€"'

        # Ajuster largeurs colonnes
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 55
        ws.column_dimensions['D'].width = 18

    # Sauvegarder
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
    workbook.save(temp_file.name)
    return temp_file.name


def create_csv_multi_files(operations_par_type, types_traites):
    """Fallback CSV si pas d'openpyxl"""
    import csv
    import zipfile

    temp_zip = tempfile.NamedTemporaryFile(delete=False, suffix='.zip')

    with zipfile.ZipFile(temp_zip.name, 'w') as zipf:
        for type_op in types_traites:
            if not operations_par_type[type_op]:
                continue

            temp_csv = tempfile.NamedTemporaryFile(
                mode='w', delete=False, suffix='.csv', newline='', encoding='utf-8')
            writer = csv.writer(temp_csv)

            # Headers
            writer.writerow([f'ANALYSE {type_op.upper()} - BANQUE POSTALE'])
            writer.writerow([])
            writer.writerow(['Date', 'Libelle', 'Description', 'Montant'])

            # Données
            total = 0
            operations_triees = sorted(
                operations_par_type[type_op], key=lambda x: x['date'])

            for op in operations_triees:
                writer.writerow([
                    op['date'].strftime('%d/%m/%Y'),
                    op['libelle'],
                    op['description'],
                    f"{float(abs(op['montant'])):.2f}"
                ])
                total += float(abs(op['montant']))

            # Total
            writer.writerow([])
            writer.writerow(['', '', 'TOTAL:', f"{total:.2f} €"])

            temp_csv.close()
            zipf.write(temp_csv.name, f'{type_op}_banque_postale.csv')
            os.unlink(temp_csv.name)

    return temp_zip.name


def create_excel_charges_fixes_with_month(selected_operations, total_charges, mois_selectionne):
    """Crée un Excel spécialisé pour les charges fixes avec indication du mois"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return create_csv_charges_fixes(selected_operations, total_charges)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Charges Fixes Mensuelles"

    # Style
    header_fill = PatternFill(start_color="28a745",
                              end_color="28a745", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))

    # En-tête principal avec période
    periode_texte = "PÉRIODE COMPLÈTE" if mois_selectionne == 'tous' else f"MOIS DE {mois_selectionne}"
    ws['A1'] = f"CHARGES FIXES MENSUELLES - {periode_texte}"
    ws['A1'].font = Font(bold=True, size=16)
    ws.merge_cells('A1:E1')

    # Informations
    ws['A2'] = f"Nombre de charges: {len(selected_operations)}"
    ws['A3'] = f"Total mensuel: {total_charges:.2f} €"
    ws['A4'] = f"Date d'analyse: {datetime.now().strftime('%d/%m/%Y à %H:%M')}"

    # Headers
    headers = ['Date', 'Type', 'Libellé', 'Description', 'Montant (€)']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    # Données triées par date
    row = 7
    operations_triees = sorted(
        selected_operations, key=lambda x: datetime.strptime(x['date'], '%d/%m/%Y'))

    for op in operations_triees:
        ws.cell(row=row, column=1, value=op['date']).border = border
        ws.cell(row=row, column=2,
                value="Prélèvement" if op['type'] == 'prelevement' else "Achat CB").border = border
        ws.cell(row=row, column=3, value=op['libelle']).border = border
        ws.cell(row=row, column=4, value=op['description']).border = border
        ws.cell(row=row, column=5, value=float(op['montant'])).border = border
        ws.cell(row=row, column=5).number_format = '#,##0.00'
        row += 1

    # Total
    row += 1
    ws.merge_cells(f'A{row}:D{row}')
    total_cell = ws.cell(
        row=row, column=1, value=f"TOTAL CHARGES FIXES - {periode_texte}")
    total_cell.font = Font(bold=True, size=14)
    total_cell.fill = PatternFill(
        start_color="FFD966", end_color="FFD966", fill_type="solid")

    total_montant = ws.cell(row=row, column=5, value=total_charges)
    total_montant.font = Font(bold=True, size=14)
    total_montant.fill = PatternFill(
        start_color="FFD966", end_color="FFD966", fill_type="solid")
    total_montant.number_format = '#,##0.00 "€"'

    # Largeurs colonnes
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 40
    ws.column_dimensions['E'].width = 15

    # Sauvegarder
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
    workbook.save(temp_file.name)
    return temp_file.name


def create_csv_charges_fixes(selected_operations, total_charges):
    """Fallback CSV pour les charges fixes"""
    import csv

    temp_csv = tempfile.NamedTemporaryFile(
        mode='w', delete=False, suffix='.csv', newline='', encoding='utf-8')
    writer = csv.writer(temp_csv)

    # Headers
    writer.writerow(['CHARGES FIXES MENSUELLES'])
    writer.writerow([])
    writer.writerow(['Date', 'Type', 'Libelle', 'Description', 'Montant'])

    # Données
    operations_triees = sorted(
        selected_operations, key=lambda x: datetime.strptime(x['date'], '%d/%m/%Y'))

    for op in operations_triees:
        writer.writerow([
            op['date'],
            "Prélèvement" if op['type'] == 'prelevement' else "Achat CB",
            op['libelle'],
            op['description'],
            f"{float(op['montant']):.2f}"
        ])

    # Total
    writer.writerow([])
    writer.writerow(['', '', '', 'TOTAL:', f"{total_charges:.2f} €"])

    temp_csv.close()
    return temp_csv.name


def calculer_statut_dossier(session_data):
    """Calcule le statut d'avancement du dossier"""
    score = 0
    if session_data.get('revenus_nets'):
        score += 25
    if session_data.get('charges_fixes'):
        score += 25
    if session_data.get('capacite_emprunt'):
        score += 25
    if session_data.get('apport'):
        score += 25

    if score >= 75:
        return "Prêt pour RDV banque"
    elif score >= 50:
        return "Dossier en cours"
    else:
        return "Données manquantes"


# VUES REFACTORISÉES - Maintenant avec templates
def accueil(request):
    """Page d'accueil. La carte « depenses mensuelles » affiche le chiffre deja
    calcule quand il existe : l'accueil devient un point de situation, pas une
    simple vitrine."""
    depenses = request.session.get('depenses_mensuelles')
    capacite = request.session.get('capacite_emprunt')
    prix_max = request.session.get('prix_achat_max')
    revenus = request.session.get('revenus_mensuels')
    return render(request, 'analyseur/accueil.html', {
        'revenus_affiches': format_euros(revenus) if revenus else None,
        'depenses_mensuelles': depenses,
        'depenses_affichees': format_euros(depenses) if depenses else None,
        'capacite_affichee': format_euros(capacite) if capacite else None,
        'prix_max_affiche': format_euros(prix_max) if prix_max else None,
        # L'étape à faire maintenant : la première pas encore faite.
        'prochaine_etape': 1 if not depenses else (2 if not capacite else 3),
    })


def upload_releve(request):
    """Interface d'upload avec parser optimisé - maintenant avec template"""
    if request.method == 'POST':
        try:
            # Récupérer les types sélectionnés
            types_selectionnes = []
            if request.POST.get('achat'):
                types_selectionnes.append('achat')
            if request.POST.get('virement'):
                types_selectionnes.append('virement')
            if request.POST.get('depot'):
                types_selectionnes.append('depot')
            if request.POST.get('debit'):
                types_selectionnes.append('debit')
            if request.POST.get('cheque'):
                types_selectionnes.append('cheque')

            if not types_selectionnes:
                raise Exception(
                    "Veuillez sélectionner au moins un type d'opération")

            # Parser le vrai PDF uploadé
            fichier_pdf = request.FILES.get('fichier_pdf')
            if not fichier_pdf:
                raise Exception("Veuillez sélectionner un fichier PDF")

            # Sauvegarder temporairement le PDF
            temp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
            for chunk in fichier_pdf.chunks():
                temp_pdf.write(chunk)
            temp_pdf.close()

            # Parser avec le nouveau parser simplifié
            parser = BanquePostaleParserSimple()
            try:
                resultats = parser.parse_pdf(temp_pdf.name, types_selectionnes)
            finally:
                if os.path.exists(temp_pdf.name):
                    os.unlink(temp_pdf.name)

            if not resultats['success']:
                raise Exception(f"Erreur de parsing: {resultats['error']}")

            if not any(resultats['operations_par_type'].values()):
                raise Exception(
                    "Aucune opération trouvée dans le PDF. Vérifiez le format du relevé.")

            # Créer Excel optimisé
            excel_path = create_excel_multi_onglets(
                resultats['operations_par_type'],
                resultats['types_traites']
            )

            # Télécharger
            content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if excel_path.endswith(
                '.xlsx') else 'application/zip'
            filename = f"analyse_banque_postale_{datetime.now().strftime('%Y%m%d_%H%M')}"
            filename += '.xlsx' if excel_path.endswith('.xlsx') else '.zip'

            try:
                with open(excel_path, 'rb') as f:
                    data = f.read()
            finally:
                if os.path.exists(excel_path):
                    os.unlink(excel_path)

            response = HttpResponse(data, content_type=content_type)
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response

        except Exception as e:
            # En cas d'erreur, afficher le template avec l'erreur
            return render(request, 'analyseur/upload_releve.html', {
                'error': str(e)
            })

    # GET - Afficher le formulaire
    return render(request, 'analyseur/upload_releve.html')


def charges_fixes(request):
    """Interface pour calculer les charges fixes mensuelles avec upload multiple et filtrage par mois"""
    if request.method == 'POST':
        # Étape 1: Upload et parsing initial (1 ou 2 fichiers)
        if 'fichier_pdf1' in request.FILES or 'fichier_pdf2' in request.FILES:
            try:
                fichiers_pdf = []

                # Traiter le premier fichier
                if 'fichier_pdf1' in request.FILES:
                    fichiers_pdf.append(request.FILES['fichier_pdf1'])

                # Traiter le second fichier (optionnel)
                if 'fichier_pdf2' in request.FILES:
                    fichiers_pdf.append(request.FILES['fichier_pdf2'])

                if not fichiers_pdf:
                    raise Exception(
                        "Veuillez sélectionner au moins un fichier PDF")

                # Parser tous les fichiers et agréger les données
                all_prelevements = []
                all_achats_cb = []
                mois_disponibles = set()

                parser = BanquePostaleParserSimple()

                for fichier_pdf in fichiers_pdf:
                    # Sauvegarder temporairement le PDF
                    temp_pdf = tempfile.NamedTemporaryFile(
                        delete=False, suffix='.pdf')
                    for chunk in fichier_pdf.chunks():
                        temp_pdf.write(chunk)
                    temp_pdf.close()

                    # Parser le fichier
                    try:
                        resultats = parser.parse_pdf(
                            temp_pdf.name, ['debit', 'achat'])
                    finally:
                        if os.path.exists(temp_pdf.name):
                            os.unlink(temp_pdf.name)

                    if not resultats['success']:
                        raise Exception(
                            f"Erreur de parsing sur {fichier_pdf.name}: {resultats['error']}")

                    # Agréger les opérations (sorties uniquement : on écarte un
                    # crédit qui aurait été rangé par mot-clé dans debit/achat).
                    prelevements = [op for op in resultats['operations_par_type'].get('debit', [])
                                    if op.get('sens') != 'credit']
                    achats_cb = [op for op in resultats['operations_par_type'].get('achat', [])
                                 if op.get('sens') != 'credit']

                    all_prelevements.extend(prelevements)
                    all_achats_cb.extend(achats_cb)

                    # Extraire les mois disponibles
                    for op in prelevements + achats_cb:
                        mois_annee = f"{op['date'].strftime('%m/%Y')}"
                        mois_nom = f"{op['date'].strftime('%B %Y')}"
                        # Conversion des noms de mois en français
                        mois_fr = mois_nom.replace('January', 'Janvier').replace('February', 'Février') \
                            .replace('March', 'Mars').replace('April', 'Avril') \
                            .replace('May', 'Mai').replace('June', 'Juin') \
                            .replace('July', 'Juillet').replace('August', 'Août') \
                            .replace('September', 'Septembre').replace('October', 'Octobre') \
                            .replace('November', 'Novembre').replace('December', 'Décembre')
                        mois_disponibles.add((mois_annee, mois_fr))

                if not all_prelevements and not all_achats_cb:
                    raise Exception(
                        "Aucun prélèvement ni achat CB trouvé dans le(s) PDF.")

                # Trier les mois chronologiquement
                mois_tries = sorted(list(mois_disponibles), key=lambda x: x[0])

                # Afficher l'interface de sélection avec choix du mois
                return render_selection_interface_with_month(all_prelevements, all_achats_cb, mois_tries, len(fichiers_pdf))

            except Exception as e:
                return render(request, 'analyseur/charges_fixes.html', {
                    'error': str(e)
                })

        # Étape 2: Traitement des sélections utilisateur avec filtrage par mois
        elif 'selected_operations' in request.POST:
            try:
                # Récupérer les opérations sélectionnées et le mois choisi
                selected_data = json.loads(request.POST['selected_operations'])
                mois_selectionne = request.POST.get('mois_filtre', 'tous')

                # Filtrer par mois si nécessaire. Les dates arrivent au format
                # jj/mm/AAAA et le mois choisi au format mm/AAAA : on compare
                # le couple mois/année, pas un préfixe (qui ne matchait jamais).
                if mois_selectionne != 'tous':
                    def _mois_annee(date_str):
                        parts = date_str.split('/')
                        return f"{parts[1]}/{parts[2]}" if len(parts) == 3 else ''
                    selected_data = [op for op in selected_data
                                     if _mois_annee(op['date']) == mois_selectionne]

                # Calculer le total
                total_charges = sum(float(op['montant'])
                                    for op in selected_data)

                # Créer l'Excel des charges fixes
                excel_path = create_excel_charges_fixes_with_month(
                    selected_data, total_charges, mois_selectionne)

                # Télécharger
                try:
                    with open(excel_path, 'rb') as f:
                        data = f.read()
                finally:
                    if os.path.exists(excel_path):
                        os.unlink(excel_path)

                filename_suffix = f"_{mois_selectionne.replace('/', '_')}" if mois_selectionne != 'tous' else '_periode_complete'
                filename = f"charges_fixes{filename_suffix}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
                response = HttpResponse(
                    data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                response['Content-Disposition'] = f'attachment; filename="{filename}"'
                return response

            except Exception as e:
                return HttpResponse(f"Erreur lors du traitement: {str(e)}")

    # GET - Formulaire initial avec upload multiple
    return render(request, 'analyseur/charges_fixes.html')


def render_selection_interface_with_month(prelevements, achats_cb, mois_disponibles, nb_fichiers):
    """Maintenant utilise un template au lieu de HTML en dur"""

    # Préparer les données pour le template
    context = {
        'prelevements': prelevements,
        'achats_cb': achats_cb,
        'mois_disponibles': mois_disponibles,
        'nb_fichiers': nb_fichiers,
        'prelevements_json': json.dumps([{
            'date': op['date'].strftime('%d/%m/%Y'),
            'mois': op['date'].strftime('%m/%Y'),
            'libelle': op['libelle'],
            'description': op['description'],
            'montant': str(op['montant']),
            'type': 'prelevement'
        } for op in prelevements]),
        'achats_json': json.dumps([{
            'date': op['date'].strftime('%d/%m/%Y'),
            'mois': op['date'].strftime('%m/%Y'),
            'libelle': op['libelle'],
            'description': op['description'],
            'montant': str(op['montant']),
            'type': 'achat'
        } for op in achats_cb])
    }

    # Utiliser un template Django au lieu de HTML en dur
    html_content = render_to_string(
        'analyseur/selection_charges.html', context)
    return HttpResponse(html_content)


# Réponses du simulateur gardées dans le navigateur (session en cookie signé, 1 an),
# pour les retrouver à la visite suivante. Les taux n'en font pas partie : ils
# viennent toujours du barème du jour.
SAISIE_SIMULATEUR = ('revenus_nets', 'charges_mensuelles', 'apport', 'type_bien',
                     'primo_accedant', 'duree', 'region', 'profil', 'nb_adultes',
                     'nb_enfants', 'age')


def _memoriser_saisie(request, data, mode):
    saisie = dict(request.session.get('simulateur_saisie', {}))
    if mode == 'capacite':
        # Une case décochée n'est pas envoyée : chaque champ est réécrit.
        saisie.update({k: str(data.get(k, '') or '') for k in SAISIE_SIMULATEUR})
    else:
        saisie.update({'montant_emprunt': str(data.get('montant_emprunt', '') or ''),
                       'duree_mensualite': str(data.get('duree', '') or '')})
    saisie['mode'] = mode
    request.session['simulateur_saisie'] = saisie


def _oublier_revenus_saisis(session):
    """De nouveaux relevés font foi pour les revenus et les crédits."""
    saisie = dict(session.get('simulateur_saisie', {}))
    for k in ('revenus_nets', 'charges_mensuelles'):
        saisie.pop(k, None)
    session['simulateur_saisie'] = saisie


def effacer_dossier(request):
    """« Effacer mon dossier » : vide tout ce que le navigateur garde."""
    if request.method == 'POST':
        request.session.flush()
    return redirect('accueil')


def simulateur_pret(request):
    """Interface principale du simulateur de prêt immobilier avec dashboard intégré"""

    if request.method == 'POST':
        try:
            data = json.loads(
                request.body) if request.content_type == 'application/json' else request.POST

            # Récupération des paramètres
            mode = data.get('mode', 'capacite')  # capacite ou mensualite

            simulateur = SimulateurPretImmobilier()

            if mode == 'capacite':
                # Mode calcul de capacité d'emprunt
                revenus = float(data.get('revenus_nets', 0))
                charges = float(data.get('charges_mensuelles', 0))
                duree = int(data.get('duree', 20))
                region = data.get('region', 'autre')
                profil = data.get('profil', 'moyen')
                age = int(data.get('age', 35))
                nb_adultes = int(data.get('nb_adultes', 2))
                nb_enfants = int(data.get('nb_enfants', 0))

                if revenus <= 0:
                    raise ValueError("Les revenus doivent être supérieurs à 0")

                resultat = simulateur.calculer_capacite_emprunt(
                    revenus, charges, duree, region, profil, age)

                # Prix d'achat maximum : l'argent disponible (capacité + apport)
                # doit couvrir le prix DU BIEN *et* les frais de notaire. Comme
                # frais = prix × taux, on inverse : prix = dispo / (1 + taux),
                # sinon les frais ne seraient réservés par rien.
                apport = float(data.get('apport', 0))
                type_bien = data.get('type_bien', 'ancien')
                primo = data.get('primo_accedant') in ('on', 'true', '1', True)
                cle_notaire = 'ancien_primo' if (primo and type_bien == 'ancien') else type_bien
                taux_notaire = FRAIS_NOTAIRE.get(cle_notaire, FRAIS_NOTAIRE['ancien'])

                budget_disponible = resultat['capacite_emprunt'] + apport
                prix_max = budget_disponible / (1 + taux_notaire)
                frais_notaire = simulateur.calculer_frais_notaire(
                    prix_max, cle_notaire)

                # Calcul reste à vivre avec alertes
                reste_vivre_data = simulateur.calculer_reste_a_vivre(
                    revenus, charges + resultat['mensualite_max'], nb_adultes, nb_enfants)

                # Sauvegarde en session pour le dashboard
                request.session['revenus_nets'] = revenus
                request.session['charges_fixes'] = charges
                request.session['capacite_emprunt'] = resultat['capacite_emprunt']
                request.session['prix_achat_max'] = round(prix_max)
                request.session['mensualite_max'] = resultat['mensualite_max']
                request.session['apport'] = apport
                request.session['duree'] = duree
                request.session['taux_nominal'] = resultat['taux_nominal']
                request.session['taux_assurance'] = resultat['taux_assurance']
                request.session['bareme_date'] = TAUX_DATE.isoformat()
                request.session['nb_adultes'] = nb_adultes
                request.session['nb_enfants'] = nb_enfants
                request.session['primo_accedant'] = primo
                # Une nouvelle simulation fait foi sur la page « Biens finançables ».
                request.session.pop('profil_biens', None)

                resultat.update({
                    'prix_achat_max': round(prix_max, 2),
                    'apport': apport,
                    'frais_notaire': frais_notaire,
                    'budget_total': round(prix_max + frais_notaire, 2),
                    'reste_a_vivre': reste_vivre_data['reste_a_vivre'],
                    'alerte_reste_vivre': reste_vivre_data['alerte'],
                    'statut_dossier': calculer_statut_dossier(request.session)
                })

            else:  # mode == 'mensualite'
                # Mode calcul de mensualités
                montant = float(data.get('montant_emprunt', 0))
                duree = int(data.get('duree', 20))
                taux_nominal = float(data.get('taux_nominal', 3.5))
                taux_assurance = float(data.get('taux_assurance', 0.25))

                if montant <= 0:
                    raise ValueError(
                        "Le montant d'emprunt doit être supérieur à 0")

                resultat = simulateur.calculer_mensualites(
                    montant, duree, taux_nominal, taux_assurance)

                # Ajout du tableau d'amortissement
                resultat['tableau_amortissement'] = simulateur.generer_tableau_amortissement(
                    montant, duree, taux_nominal, taux_assurance
                )

            _memoriser_saisie(request, data, mode)
            return JsonResponse({'success': True, 'data': resultat})

        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})

    # GET - Interface utilisateur avec template Django. Si les depenses
    # mensuelles ont deja ete analysees, on prerempli le champ « charges » :
    # l'utilisateur n'a pas a recopier un chiffre que l'outil connait deja.
    return render(request, 'analyseur/simulateur_pret.html', {
        'saisie': request.session.get('simulateur_saisie', {}),
        'charges_credits': request.session.get('charges_credits'),
        'loyer_actuel': request.session.get('loyer_actuel'),
        'revenus_mensuels': request.session.get('revenus_mensuels'),
    })


def dashboard_dossier(request):
    """Dashboard centralisé du dossier immobilier"""

    # Récupérer les données utilisateur (session)
    revenus = float(request.session.get('revenus_nets', 0) or 0)
    charges = float(request.session.get('charges_fixes', 0) or 0)
    mensualite_max = float(request.session.get('mensualite_max', 0) or 0)

    # Le taux d'endettement est calculé ici : les templates Django n'ont pas
    # d'opérateurs arithmétiques (les filtres « div »/« mul » n'existent pas).
    taux_endettement = round(
        (charges + mensualite_max) / revenus * 100, 1) if revenus > 0 else 0

    dossier_data = {
        'revenus_nets': revenus,
        'charges_fixes': charges,
        'capacite_calculee': request.session.get('capacite_emprunt', 0),
        'apport': request.session.get('apport', 0),
        'mensualite_max': mensualite_max,
        'taux_endettement': taux_endettement,
        'statut_dossier': calculer_statut_dossier(request.session),
    }

    return render(request, 'analyseur/dashboard.html', {'dossier': dossier_data})


def export_dossier_pdf(request):
    """Génère un PDF synthétique du dossier"""
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4
    except ImportError:
        return HttpResponse("Module reportlab non installé", status=500)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="dossier_immobilier_{datetime.now().strftime("%Y%m%d")}.pdf"'

    p = canvas.Canvas(response, pagesize=A4)

    # En-tête
    p.setFont("Helvetica-Bold", 16)
    p.drawString(50, 800, "DOSSIER DE FINANCEMENT IMMOBILIER")
    p.setFont("Helvetica", 10)
    p.drawString(50, 780, f"Généré le {datetime.now().strftime('%d/%m/%Y')}")

    # Données financières
    y_pos = 750
    session_data = request.session

    sections = [
        ("SITUATION FINANCIÈRE", [
            f"Revenus nets mensuels: {session_data.get('revenus_nets', 'N/A')} €",
            f"Charges fixes mensuelles: {session_data.get('charges_fixes', 'N/A')} €",
            f"Apport personnel: {session_data.get('apport', 'N/A')} €"
        ]),
        ("CAPACITÉ D'EMPRUNT", [
            f"Montant maximum: {session_data.get('capacite_emprunt', 'N/A')} €",
            f"Mensualité maximum: {session_data.get('mensualite_max', 'N/A')} €",
            f"Durée recommandée: {session_data.get('duree', 'N/A')} ans"
        ])
    ]

    for titre, items in sections:
        p.setFont("Helvetica-Bold", 12)
        p.drawString(50, y_pos, titre)
        y_pos -= 20

        p.setFont("Helvetica", 10)
        for item in items:
            p.drawString(70, y_pos, f"• {item}")
            y_pos -= 15
        y_pos -= 10

    p.save()
    return response


# ---------------------------------------------------------------------------
# Tableau de bord des flux du compte : ce qui sort, ce qui rentre, mois par mois
#
# L'analyse ne porte que sur des mois civils entièrement couverts par les
# relevés fournis. Un relevé Banque Postale court du 12 d'un mois au 11 du
# suivant : aucun relevé ne contient donc un mois civil à lui seul, mais deux
# relevés consécutifs se complètent. N relevés qui se suivent = N-1 mois
# complets (4 relevés → 3 mois, 6 relevés → 5 mois).
# ---------------------------------------------------------------------------

MOIS_FR = {
    1: 'Janvier', 2: 'Février', 3: 'Mars', 4: 'Avril', 5: 'Mai', 6: 'Juin',
    7: 'Juillet', 8: 'Août', 9: 'Septembre', 10: 'Octobre', 11: 'Novembre',
    12: 'Décembre',
}

MOIS_COURT_FR = {
    1: 'janv.', 2: 'févr.', 3: 'mars', 4: 'avril', 5: 'mai', 6: 'juin',
    7: 'juil.', 8: 'août', 9: 'sept.', 10: 'oct.', 11: 'nov.', 12: 'déc.',
}

# Libellés des catégories du parser, tels qu'ils parlent à l'utilisateur.
CATEGORIES_SORTIES = {
    'achat': 'Achats carte (courses, essence, quotidien)',
    'debit': 'Prélèvements (abonnements, énergie, impôts)',
    'virement': 'Virements sortants',
    'cheque': 'Chèques',
    'autre': 'Autres sorties',
    'depot': 'Divers',
}

CATEGORIES_ENTREES = {
    'virement': 'Virements reçus (salaires, remboursements)',
    'depot': 'Dépôts et remises',
    'achat': 'Remboursements carte',
    'debit': 'Régularisations',
    'cheque': 'Chèques encaissés',
    'autre': 'Autres entrées',
}

# Un poste présent sur au moins cette part des mois analysés est considéré
# comme récurrent. Ajustable à l'écran : le bon seuil dépend du nombre de mois.
SEUIL_RECURRENCE_PAR_DEFAUT = 40


def format_euros(valeur):
    """Montant arrondi à l'euro, séparateur de milliers français.

    Les templates Django n'ont pas de filtre pour ça (`intcomma` de humanize
    produit une virgule anglo-saxonne, et l'app n'installe pas humanize) : le
    formatage se fait donc ici, comme tous les autres calculs d'affichage.
    """
    return '{:,.0f}'.format(round(float(valeur))).replace(',', ' ')


def _dernier_jour_du_mois(annee, mois):
    if mois == 12:
        return date(annee, 12, 31)
    return date(annee, mois + 1, 1) - timedelta(days=1)


# Un relevé mensuel manquant creuse un écart d'environ 30 jours. En deçà de ce
# seuil, l'écart entre deux relevés qui se suivent ne peut pas cacher un relevé
# entier : ce sont des jours sans mouvement, pas une absence de couverture.
ECART_MAX_SANS_OPERATION = 20


def _fusionner_periodes(periodes):
    """Fusionne les périodes des relevés en segments de couverture continue.

    Deux relevés qui se suivent (le premier finit le 11, le second commence le
    12) forment un seul segment : c'est ce qui permet à un mois civil d'être
    couvert par deux relevés différents.

    Les bornes d'un relevé sont parfois déduites de ses opérations, faute d'en
    lire la période dans le PDF. Un week-end sans dépense en fin de relevé et
    trois jours calmes au début du suivant créaient alors un faux trou, qui
    suffisait à écarter le mois entier de l'analyse. On comble donc les écarts
    trop courts pour contenir un relevé manquant.
    """
    segments = []
    for debut, fin in sorted(periodes):
        if segments and debut <= segments[-1][1] + timedelta(
                days=ECART_MAX_SANS_OPERATION):
            segments[-1][1] = max(segments[-1][1], fin)
        else:
            segments.append([debut, fin])
    return [(d, f) for d, f in segments]


def _mois_entierement_couvert(annee, mois, segments):
    premier = date(annee, mois, 1)
    dernier = _dernier_jour_du_mois(annee, mois)
    return any(d <= premier and f >= dernier for d, f in segments)


def _dedupliquer(operations, periodes):
    """Retire les opérations comptées deux fois par des relevés qui se
    recouvrent.

    La déduplication ne s'applique QUE dans les zones couvertes par plusieurs
    relevés : ailleurs, deux opérations identiques le même jour (deux passages
    à la même station-service) sont bien deux dépenses distinctes et doivent
    rester.
    """
    zones = [
        (max(a[0], b[0]), min(a[1], b[1]))
        for i, a in enumerate(periodes) for b in periodes[i + 1:]
        if max(a[0], b[0]) <= min(a[1], b[1])
    ]
    if not zones:
        return operations, 0

    retenues, vues, retirees = [], set(), 0
    for op in operations:
        if any(d <= op['date'] <= f for d, f in zones):
            cle = (op['date'], op['montant'], op['description'], op.get('sens'))
            if cle in vues:
                retirees += 1
                continue
            vues.add(cle)
        retenues.append(op)
    return retenues, retirees


# Préfixes qui décrivent le TYPE d'opération, pas le bénéficiaire : les retirer
# fait converger « ACHAT CB CARREFOUR » et « CARREFOUR » vers le même poste.
_PREFIXES_OPERATION = re.compile(
    r'^(?:ACHAT\s+CB|PAIEMENT\s+CB|ACHAT|PAIEMENT|PRELEVEMENT\s+DE|'
    r'PRELEVEMENT|VIREMENT\s+INSTANTANE\s*[AD]?E?|VIREMENT\s+PERMANENT|'
    r'VIREMENT\s+POUR|VIREMENT\s+DE|VIREMENT|RETRAIT|CHEQUE|VERSEMENT|REMISE)\s*',
    re.IGNORECASE)


def _cle_poste(libelle):
    """Clé de regroupement d'un libellé, pour repérer les postes récurrents.

    On retire le préfixe de type d'opération, les chiffres et la ponctuation,
    puis on ne garde que les deux premiers mots : « ACHAT CB CARREFOUR 1234 »
    et « CARREFOUR.FR » se rejoignent, sans fusionner deux commerçants
    différents.
    """
    texte = _PREFIXES_OPERATION.sub('', (libelle or '').upper())
    texte = re.sub(r'[^A-Z ]+', ' ', texte)
    mots = [m for m in texte.split() if len(m) > 2]
    return ' '.join(mots[:2]) if mots else (libelle or '').upper()[:20]


def _repartition(operations, libelles_categories):
    """Répartition par catégorie, de la plus lourde à la plus légère."""
    total = sum(op['montant'] for op in operations)
    par_categorie = {}
    for op in operations:
        cat = op.get('type', 'autre')
        par_categorie[cat] = par_categorie.get(cat, Decimal('0')) + op['montant']
    return sorted(
        ({'libelle': libelles_categories.get(cle, cle.capitalize()),
          'montant': format_euros(montant),
          'brut': float(montant),
          'part': round(float(montant) / float(total) * 100) if total else 0}
         for cle, montant in par_categorie.items()),
        key=lambda c: c['brut'], reverse=True)


def _ligne_operation(op):
    """Une opération telle qu'affichée dans les tableaux détaillés."""
    return {
        'date': op['date'].strftime('%d/%m/%Y'),
        'jour': op['date'].strftime('%d/%m'),
        'libelle': (op['libelle'] or op['description'])[:60],
        'categorie': (CATEGORIES_SORTIES if op['sens'] == 'debit'
                      else CATEGORIES_ENTREES).get(op.get('type', 'autre'),
                                                   'Autres'),
        'montant': float(op['montant']),
        'montant_affiche': format_euros(op['montant']),
        'sens': op['sens'],
        'interne': bool(op.get('interne')),
        'beneficiaire': (op.get('complement') or '')[:50],
        # un virement se reconnaît à son bénéficiaire, pas à « VIREMENT POUR »
        'poste': _cle_poste(op['complement'] if op.get('complement') and
                            re.match(r'(?i)VIR', op['libelle'] or op['description'] or '')
                            else op['libelle'] or op['description']),
    }


def _postes_recurrents(mois_complets, nb_mois):
    """Postes qui reviennent d'un mois sur l'autre, entrants comme sortants.

    Un poste est décrit par sa clé normalisée. On retient sa fréquence (nombre
    de mois où il apparaît) et son montant mensuel moyen — c'est cette base qui
    permet de distinguer une charge fixe d'une dépense ponctuelle.
    """
    postes = {}
    for m in mois_complets:
        vus_ce_mois = {}
        for op in m['operations']:
            if op.get('interne'):
                continue    # l'épargne n'est pas une charge fixe
            cle = (op['poste'], op['sens'])
            entree = vus_ce_mois.setdefault(cle, {'montant': 0.0, 'nb': 0})
            entree['montant'] += op['montant']
            entree['nb'] += 1
        for cle, agrege in vus_ce_mois.items():
            poste = postes.setdefault(cle, {
                'poste': cle[0], 'sens': cle[1],
                'mois': [], 'montants': [], 'nb_operations': 0,
                'exemple': '',
            })
            poste['mois'].append(m['court'])
            poste['montants'].append(agrege['montant'])
            poste['nb_operations'] += agrege['nb']
            if not poste['exemple']:
                poste['exemple'] = next(
                    op['libelle'] for op in m['operations']
                    if op['poste'] == cle[0] and op['sens'] == cle[1])

    resultat = []
    for poste in postes.values():
        nb = len(poste['mois'])
        moyen = sum(poste['montants']) / nb
        resultat.append({
            'poste': poste['poste'],
            'exemple': poste['exemple'],
            'sens': poste['sens'],
            'nb_mois': nb,
            'frequence': round(nb / nb_mois * 100),
            'mois': ', '.join(poste['mois']),
            'montant_moyen': round(moyen, 2),
            'montant_moyen_affiche': format_euros(moyen),
            'montant_total': round(sum(poste['montants']), 2),
            'montant_total_affiche': format_euros(sum(poste['montants'])),
            'nb_operations': poste['nb_operations'],
        })
    # Les plus fréquents d'abord, puis les plus lourds : c'est l'ordre dans
    # lequel on veut lire ses charges fixes.
    resultat.sort(key=lambda p: (-p['frequence'], -p['montant_moyen']))
    return resultat


# Ce que la banque compte dans les 35 % d'endettement : les crédits en cours
# (et pensions versées), PAS les dépenses courantes. Le loyer actuel disparaît
# avec l'achat de la résidence principale : il est rendu à part, pour le
# « saut de charge », jamais ajouté aux charges.
_RE_CREDIT = re.compile(
    r"\b(PRET|ECHEANCE|ECH PRET|CREDIT IMMO|CREDIT CONSO|COFIDIS|CETELEM|SOFINCO|"
    r"FRANFINANCE|ONEY|FLOA|YOUNITED|COFINOGA|FINANCO|DIAC|RCI BANQUE|PSA BANQUE|"
    r"CA CONSUMER|BNP PERSONAL|LOA|LLD|PENSION ALIMENTAIRE)\b")
_RE_LOYER = re.compile(
    r"\b(LOYER|LOYERS|BAILLEUR|FONCIA|NEXITY|CITYA|ORALIA|PARIS HABITAT|OPH|HLM|"
    r"VALOPHIS|SEQENS|ICF HABITAT|CDC HABITAT|IN LI|IMMOBILIERE 3F)\b")
FREQUENCE_MINI_CHARGE = 50   # % des mois : en dessous, c'est ponctuel, pas une mensualité


def _sans_accents_maj(texte):
    return unicodedata.normalize('NFKD', texte or '').encode('ascii', 'ignore').decode().upper()


_RE_EPARGNE = re.compile(r"\b(LIVRET|LDDS|LDD|LEP|PEL|CEL|EPARGNE|ECONOMIES?|ASSURANCE VIE|PERP)\b")


def _proche(a, b):
    """Même nom à une faute de frappe près (DUPONT / DUPPONT)."""
    if abs(len(a) - len(b)) > 1:
        return False
    if len(a) == len(b):
        return sum(x != y for x, y in zip(a, b)) <= 1
    court, long_ = sorted((a, b), key=len)
    return any(long_[:k] + long_[k + 1:] == court for k in range(len(long_)))


def est_virement_interne(op, titulaires):
    """Virement entre comptes du foyer : vers/depuis un titulaire, ou vers l'épargne.

    Ce n'est ni une dépense ni un revenu (CC → PEL, LDDS, autre banque à soi) :
    l'argent reste au foyer. Un remboursement de crédit n'est jamais interne.
    """
    texte = _sans_accents_maj(f"{op.get('description', '')} {op.get('complement', '')}")
    if not re.search(r"\bVIR", texte) or _RE_CREDIT.search(texte):
        return False
    if _RE_EPARGNE.search(texte):
        return True
    mots = [m for m in re.findall(r"[A-Z]{5,}", texte)]
    return any(_proche(m, t) for m in mots for t in titulaires if len(t) >= 5)


def charges_bancaires(postes, nb_mois):
    """Mensualités de crédit et loyer repérés dans les postes récurrents sortants.

    Moyenne sur TOUS les mois complets (montant total ÷ nb de mois) : un
    prélèvement vu 2 mois sur 3 pèse pour ce qu'il coûte réellement par mois.
    """
    def montant(p):
        return round(p['montant_total'] / nb_mois, 2) if nb_mois else 0.0

    credits, loyers = [], []
    for p in postes:
        if p['sens'] != 'debit' or p['frequence'] < FREQUENCE_MINI_CHARGE:
            continue
        texte = _sans_accents_maj(f"{p['exemple']} {p['poste']}")
        cible = credits if _RE_CREDIT.search(texte) else loyers if _RE_LOYER.search(texte) else None
        if cible is not None:
            cible.append({'exemple': p['exemple'], 'mensuel': montant(p),
                          'mensuel_affiche': format_euros(montant(p))})
    credits.sort(key=lambda c: -c['mensuel'])
    loyers.sort(key=lambda c: -c['mensuel'])
    return {'credits': credits, 'credits_mensuels': round(sum(c['mensuel'] for c in credits), 2),
            'loyers': loyers, 'loyer_mensuel': round(sum(c['mensuel'] for c in loyers), 2)}


def analyser_flux_mensuels(releves):
    """Analyse complète des flux du compte, par mois civil entièrement couvert.

    `releves` : liste de dicts {nom, periode, operations}. Débits ET crédits
    sont conservés — le sens vient de la colonne du PDF, seule source fiable.

    Renvoie None si aucun mois n'est complet : mieux vaut expliquer ce qui
    manque que présenter un total amputé de la moitié d'un mois.
    """
    # `detail` retrace ce que chaque fichier a apporté. C'est le seul moyen,
    # face à un trou de couverture, de distinguer un relevé réellement absent
    # d'un fichier mal lu.
    periodes, operations, detail, ecarts_controle = [], [], [], []
    for releve in releves:
        connues = [op for op in releve['operations']
                   if op.get('sens') in ('debit', 'credit')]
        if not connues:
            detail.append({'nom': releve['nom'], 'periode': 'aucune opération lue',
                           'debut': date.max, 'nb_operations': 0,
                           'source': 'illisible',
                           'controle': {'statut': 'absent', 'ecart': ''}})
            continue
        operations.extend(connues)
        dates = [op['date'] for op in connues]
        lue = releve.get('periode')
        debut, fin = lue or (min(dates), max(dates))
        debut, fin = min(debut, min(dates)), max(fin, max(dates))
        periodes.append((debut, fin))
        # Contrôle : ce qu'on a additionné doit égaler ce que le relevé
        # annonce. Un écart signale un montant mal lu, donc une analyse à
        # ne pas présenter à une banque.
        totaux = releve.get('totaux_releve')
        controle = {'statut': 'absent', 'ecart': ''}
        if totaux:
            attendu_d, attendu_c = float(totaux[0]), float(totaux[1])
            lu_d = float(sum(op['montant'] for op in connues
                             if op['sens'] == 'debit'))
            lu_c = float(sum(op['montant'] for op in connues
                             if op['sens'] == 'credit'))
            ecart = max(abs(attendu_d - lu_d), abs(attendu_c - lu_c))
            controle = {
                'statut': 'ok' if ecart < 1 else 'ecart',
                'ecart': format_euros(ecart) if ecart >= 1 else '',
                'attendu': f"{format_euros(attendu_d)} / {format_euros(attendu_c)}",
            }
            if ecart >= 1:
                ecarts_controle.append(f"{releve['nom']} ({format_euros(ecart)} €)")

        detail.append({
            'nom': releve['nom'],
            'periode': f"{debut:%d/%m/%Y} → {fin:%d/%m/%Y}",
            'debut': debut,
            'nb_operations': len(connues),
            'source': 'lue dans le PDF' if lue else 'déduite des opérations',
            'controle': controle,
        })

    # Chronologique ; les fichiers muets, sans date exploitable, ferment la
    # liste (date.max) plutôt que de la parasiter en tête.
    detail.sort(key=lambda d: d['debut'])

    if not operations:
        return None

    operations, doublons = _dedupliquer(operations, periodes)
    segments = _fusionner_periodes(periodes)

    # Virements entre comptes du foyer : repérés APRÈS le contrôle des totaux
    # (qui doit voir toutes les opérations du relevé), puis tenus à l'écart
    # des dépenses et des revenus.
    titulaires = sorted({t for r in releves for t in r.get('titulaires', [])})
    for op in operations:
        op['interne'] = est_virement_interne(op, titulaires)

    par_mois = {}
    for op in operations:
        par_mois.setdefault((op['date'].year, op['date'].month), []).append(op)

    mois_complets, mois_ecartes = [], []
    for (annee, mois) in sorted(par_mois):
        libelle = f"{MOIS_FR[mois]} {annee}"
        if not _mois_entierement_couvert(annee, mois, segments):
            mois_ecartes.append(libelle)
            continue

        du_mois = par_mois[(annee, mois)]
        debits = [op for op in du_mois if op['sens'] == 'debit' and not op.get('interne')]
        credits = [op for op in du_mois if op['sens'] == 'credit' and not op.get('interne')]
        sorties = float(sum(op['montant'] for op in debits))
        entrees = float(sum(op['montant'] for op in credits))
        internes_sortants = float(sum(op['montant'] for op in du_mois
                                      if op['sens'] == 'debit' and op.get('interne')))
        internes_entrants = float(sum(op['montant'] for op in du_mois
                                      if op['sens'] == 'credit' and op.get('interne')))

        lignes = sorted((_ligne_operation(op) for op in du_mois),
                        key=lambda l: (-l['montant'],))
        mois_complets.append({
            'cle': f"{annee}-{mois:02d}",
            'libelle': libelle,
            'court': f"{MOIS_COURT_FR[mois]} {annee}",
            'sorties': round(sorties, 2),
            'sorties_affichees': format_euros(sorties),
            'entrees': round(entrees, 2),
            'entrees_affichees': format_euros(entrees),
            'solde': round(entrees - sorties, 2),
            'solde_affiche': format_euros(abs(entrees - sorties)),
            'solde_positif': entrees >= sorties,
            'taux_effort': round(sorties / entrees * 100) if entrees else 0,
            'internes_sortants': round(internes_sortants, 2),
            'internes_entrants': round(internes_entrants, 2),
            'nb_sorties': len(debits),
            'nb_entrees': len(credits),
            'categories_sorties': _repartition(debits, CATEGORIES_SORTIES),
            'categories_entrees': _repartition(credits, CATEGORIES_ENTREES),
            'operations': lignes,
        })

    if not mois_complets:
        return None

    # Trou de couverture : un relevé manquant au milieu de la série coupe la
    # continuité, et des mois disparaissent silencieusement de l'analyse.
    trous = [
        f"{segments[i][1] + timedelta(days=1):%d/%m/%Y} au "
        f"{segments[i + 1][0] - timedelta(days=1):%d/%m/%Y}"
        for i in range(len(segments) - 1)
    ]

    nb = len(mois_complets)
    sorties_moy = sum(m['sorties'] for m in mois_complets) / nb
    entrees_moy = sum(m['entrees'] for m in mois_complets) / nb
    internes_s_moy = sum(m['internes_sortants'] for m in mois_complets) / nb
    internes_e_moy = sum(m['internes_entrants'] for m in mois_complets) / nb

    # Échelle commune aux deux séries du graphique : sans elle, entrées et
    # sorties ne seraient pas comparables d'un coup d'œil.
    maximum = max(max(m['sorties'], m['entrees']) for m in mois_complets) or 1
    for m in mois_complets:
        m['hauteur_sorties'] = round(m['sorties'] / maximum * 100)
        m['hauteur_entrees'] = round(m['entrees'] / maximum * 100)

    postes = _postes_recurrents(mois_complets, nb)
    tendance = None
    if nb >= 2:
        dernier = mois_complets[-1]
        reference = sum(m['sorties'] for m in mois_complets[:-1]) / (nb - 1)
        ecart = dernier['sorties'] - reference
        variation = (ecart / reference * 100) if reference else 0
        tendance = {
            'sens': 'stable' if abs(variation) < 5 else (
                'hausse' if variation > 0 else 'baisse'),
            'variation': round(abs(variation)),
            'ecart': format_euros(abs(ecart)),
            'mois': dernier['libelle'],
        }

    return {
        'mois': mois_complets,
        'nb_mois': nb,
        'premier_mois': mois_complets[0]['libelle'],
        'dernier_mois': mois_complets[-1]['libelle'],
        'postes_recurrents': postes,
        'charges_bancaires': charges_bancaires(postes, nb),
        'seuil_recurrence': SEUIL_RECURRENCE_PAR_DEFAUT,
        'resume': {
            'entrees_moyennes': round(entrees_moy, 2),
            'entrees_moyennes_affichees': format_euros(entrees_moy),
            'sorties_moyennes': round(sorties_moy, 2),
            'sorties_moyennes_affichees': format_euros(sorties_moy),
            'internes_sortants_moyens': round(internes_s_moy, 2),
            'internes_sortants_affiches': format_euros(internes_s_moy),
            'internes_entrants_moyens': round(internes_e_moy, 2),
            'internes_entrants_affiches': format_euros(internes_e_moy),
            'titulaires': titulaires,
            'reste': round(entrees_moy - sorties_moy, 2),
            'reste_affiche': format_euros(abs(entrees_moy - sorties_moy)),
            'reste_positif': entrees_moy >= sorties_moy,
            'taux_effort': round(sorties_moy / entrees_moy * 100) if entrees_moy else 0,
            'total_entrees_affiche': format_euros(
                sum(m['entrees'] for m in mois_complets)),
            'total_sorties_affiche': format_euros(
                sum(m['sorties'] for m in mois_complets)),
        },
        'tendance': tendance,
        'detail_releves': detail,
        'releves_muets': [d['nom'] for d in detail if d['source'] == 'illisible'],
        'mois_ecartes': mois_ecartes,
        'trous': trous,
        'doublons_retires': doublons,
        'ecarts_controle': ecarts_controle,
        'controle_ok': not ecarts_controle and any(
            d['controle']['statut'] == 'ok' for d in detail),
        'couverture_debut': min(d for d, _ in segments).strftime('%d/%m/%Y'),
        'couverture_fin': max(f for _, f in segments).strftime('%d/%m/%Y'),
        'nb_operations': sum(m['nb_sorties'] + m['nb_entrees']
                             for m in mois_complets),
    }


def _lire_releves(fichiers):
    """Parse les PDF uploadés. Lève une exception au premier fichier illisible
    en tant que PDF ; un fichier lisible mais sans opération est laissé passer,
    l'analyse le signalera comme muet."""
    parser = BanquePostaleParserSimple()
    releves = []
    for fichier in fichiers:
        temp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
        try:
            for chunk in fichier.chunks():
                temp_pdf.write(chunk)
            temp_pdf.close()
            resultats = parser.parse_pdf(temp_pdf.name)
        finally:
            if os.path.exists(temp_pdf.name):
                os.unlink(temp_pdf.name)

        if not resultats['success']:
            raise Exception(f"{fichier.name} : {resultats['error']}")

        operations = []
        for type_op, ops in resultats['operations_par_type'].items():
            for op in ops:
                op['type'] = type_op
                operations.append(op)

        releves.append({
            'nom': fichier.name,
            'periode': resultats.get('periode'),
            'totaux_releve': resultats.get('totaux_releve'),
            'titulaires': resultats.get('titulaires', []),
            'operations': operations,
        })
    return releves


def depenses_mensuelles(request):
    """Tableau de bord des flux du compte. Rien n'est stocké : seules les
    moyennes sont mémorisées en session, pour préremplir le simulateur."""
    if request.method != 'POST':
        return render(request, 'analyseur/depenses_mensuelles.html')

    fichiers = request.FILES.getlist('fichiers_pdf')
    if not fichiers:
        return render(request, 'analyseur/depenses_mensuelles.html', {
            'error': "Ajoutez vos relevés PDF. Il en faut au moins deux qui se "
                     "suivent pour reconstituer un mois entier."
        })

    try:
        releves = _lire_releves(fichiers)
    except Exception as e:
        return render(request, 'analyseur/depenses_mensuelles.html',
                      {'error': str(e)})

    toutes = [op for r in releves for op in r['operations']]
    if not toutes:
        return render(request, 'analyseur/depenses_mensuelles.html', {
            'error': "Aucune opération lue dans ce(s) PDF. Vérifiez qu'il "
                     "s'agit bien d'un relevé de compte Banque Postale."
        })

    # Sans pdfplumber le sens débit/crédit est inconnu : mieux vaut le dire que
    # d'afficher un total silencieusement faux.
    if all(op.get('sens') is None for op in toutes):
        return render(request, 'analyseur/depenses_mensuelles.html', {
            'error': "Impossible de distinguer les débits des crédits sur ce "
                     "relevé (module pdfplumber indisponible). Installez les "
                     "dépendances avec « pip install -r requirements.txt »."
        })

    analyse = analyser_flux_mensuels(releves)
    if not analyse:
        # Cas courant du premier essai : un seul relevé, qui court du 12 au 11
        # et ne contient donc aucun mois civil entier.
        dates = [op['date'] for op in toutes if op.get('sens')]
        couverture = (f" Vos relevés couvrent du {min(dates):%d/%m/%Y} au "
                      f"{max(dates):%d/%m/%Y}.") if dates else ''
        return render(request, 'analyseur/depenses_mensuelles.html', {
            'error': "Aucun mois entier dans ces relevés." + couverture +
                     " Un relevé va du 12 d'un mois au 11 du suivant : il en "
                     "faut deux qui se suivent pour reconstituer un mois "
                     "complet, et N+1 pour analyser N mois."
        })

    # Report vers le simulateur. Les « charges » de l'endettement ne sont que
    # les crédits en cours : toutes les sorties y faisaient passer n'importe
    # quel bien « hors budget ». Les dépenses totales restent pour l'accueil.
    request.session['depenses_mensuelles'] = analyse['resume']['sorties_moyennes']
    request.session['revenus_mensuels'] = analyse['resume']['entrees_moyennes']
    request.session['charges_credits'] = analyse['charges_bancaires']['credits_mensuels']
    request.session['loyer_actuel'] = analyse['charges_bancaires']['loyer_mensuel']
    _oublier_revenus_saisis(request.session)

    return render(request, 'analyseur/depenses_mensuelles.html', {
        'analyse': analyse,
        'nb_fichiers': len(fichiers),
        # Embarqué dans la page : l'export Excel le renvoie tel quel, augmenté
        # des lignes saisies à la main. L'app n'ayant pas de base, c'est la
        # page elle-même qui porte l'état entre l'analyse et l'export.
        'analyse_json': json.dumps(analyse, ensure_ascii=False, default=str),
    })


# ---------------------------------------------------------------------------
# Export Excel du tableau de bord
# ---------------------------------------------------------------------------

def _feuille(classeur, titre, entetes, lignes, largeurs=None):
    """Onglet standard : en-tête bleue figée, colonnes dimensionnées.

    openpyxl est importé ici et non en tête de module : le reste du fichier le
    traite comme optionnel (repli CSV de `create_excel_multi_onglets`), et un
    import global ferait échouer l'application entière s'il manquait.
    """
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    feuille = classeur.create_sheet(titre)
    feuille.append(entetes)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color='FFFFFF')
        cellule.fill = PatternFill('solid', fgColor='2549B5')
        cellule.alignment = Alignment(horizontal='center')
    for ligne in lignes:
        feuille.append(ligne)
    for i, largeur in enumerate(largeurs or [], start=1):
        feuille.column_dimensions[get_column_letter(i)].width = largeur
    feuille.freeze_panes = 'A2'
    return feuille


def export_depenses_excel(request):
    """Classeur reprenant ce que la page affiche, lignes ajoutées comprises.

    L'application ne persiste rien : c'est la page qui renvoie l'analyse dans
    un champ caché. L'export est donc sans état, et reflète exactement ce que
    l'utilisateur avait sous les yeux — y compris ses ajustements manuels.
    """
    if request.method != 'POST':
        return redirect('depenses_mensuelles')

    try:
        from openpyxl import Workbook
    except ImportError:
        return HttpResponse("Module openpyxl non installé", status=500)

    try:
        analyse = json.loads(request.POST.get('analyse_json') or '{}')
        manuelles = json.loads(request.POST.get('lignes_manuelles') or '[]')
    except json.JSONDecodeError:
        return HttpResponse("Données d'analyse illisibles", status=400)
    if not analyse.get('mois'):
        return redirect('depenses_mensuelles')

    resume = analyse['resume']
    sorties_manuelles = sum(float(l['montant']) for l in manuelles
                            if l.get('sens') == 'debit')
    entrees_manuelles = sum(float(l['montant']) for l in manuelles
                            if l.get('sens') == 'credit')
    sorties_totales = resume['sorties_moyennes'] + sorties_manuelles
    entrees_totales = resume['entrees_moyennes'] + entrees_manuelles

    classeur = Workbook()
    classeur.remove(classeur.active)

    # --- Synthèse : ce qu'un conseiller regarde en premier ------------------
    lignes = [
        ['Période analysée',
         f"{analyse['premier_mois']} → {analyse['dernier_mois']}"],
        ['Mois entiers analysés', analyse['nb_mois']],
        ['Relevés fournis', len(analyse.get('detail_releves', []))],
        ['Contrôle des totaux',
         'Concordant avec les relevés' if analyse.get('controle_ok')
         else 'ÉCART DÉTECTÉ — vérifier'],
        [],
        ['ENTRÉES moyennes (€/mois)', round(resume['entrees_moyennes'], 2)],
        ['SORTIES moyennes (€/mois)', round(resume['sorties_moyennes'], 2)],
        ['Reste à vivre moyen (€/mois)', round(resume['reste'], 2)],
        ["Taux d'effort (sorties / entrées)", f"{resume['taux_effort']} %"],
    ]
    if manuelles:
        lignes += [
            [],
            ['Charges ajoutées à la main (€/mois)', round(sorties_manuelles, 2)],
            ['Revenus ajoutés à la main (€/mois)', round(entrees_manuelles, 2)],
            ['SORTIES corrigées (€/mois)', round(sorties_totales, 2)],
            ['ENTRÉES corrigées (€/mois)', round(entrees_totales, 2)],
            ['Reste à vivre corrigé (€/mois)',
             round(entrees_totales - sorties_totales, 2)],
            ["Taux d'effort corrigé",
             f"{round(sorties_totales / entrees_totales * 100) if entrees_totales else 0} %"],
        ]
    _feuille(classeur, 'Synthèse', ['Indicateur', 'Valeur'], lignes, [38, 34])

    # --- Mois par mois ------------------------------------------------------
    _feuille(
        classeur, 'Mois par mois',
        ['Mois', 'Entrées (€)', 'Sorties (€)', 'Solde (€)',
         "Taux d'effort", 'Nb entrées', 'Nb sorties'],
        [[m['libelle'], m['entrees'], m['sorties'], m['solde'],
          f"{m['taux_effort']} %", m['nb_entrees'], m['nb_sorties']]
         for m in analyse['mois']],
        [18, 14, 14, 14, 14, 12, 12])

    # --- Postes récurrents --------------------------------------------------
    seuil = int(request.POST.get('seuil_recurrence') or
                analyse.get('seuil_recurrence', SEUIL_RECURRENCE_PAR_DEFAUT))
    _feuille(
        classeur, 'Postes récurrents',
        ['Poste', 'Sens', 'Fréquence', 'Mois concernés',
         'Montant moyen (€/mois)', 'Total période (€)', 'Nb opérations'],
        [[p['exemple'], 'Sortie' if p['sens'] == 'debit' else 'Entrée',
          f"{p['frequence']} %", p['mois'], p['montant_moyen'],
          p['montant_total'], p['nb_operations']]
         for p in analyse['postes_recurrents'] if p['frequence'] >= seuil],
        [36, 10, 12, 30, 22, 20, 14])

    # --- Toutes les opérations ---------------------------------------------
    _feuille(
        classeur, 'Opérations',
        ['Mois', 'Date', 'Libellé', 'Catégorie', 'Sens', 'Montant (€)'],
        [[m['libelle'], op['date'], op['libelle'], op['categorie'],
          'Sortie' if op['sens'] == 'debit' else 'Entrée', op['montant']]
         for m in analyse['mois'] for op in m['operations']],
        [16, 12, 46, 34, 10, 14])

    # --- Lignes ajoutées à la main -----------------------------------------
    if manuelles:
        _feuille(
            classeur, 'Lignes ajoutées',
            ['Libellé', 'Sens', 'Montant mensuel (€)'],
            [[l.get('libelle', ''),
              'Sortie' if l.get('sens') == 'debit' else 'Entrée',
              float(l['montant'])] for l in manuelles],
            [40, 12, 22])

    # --- Couverture des relevés --------------------------------------------
    _feuille(
        classeur, 'Relevés',
        ['Fichier', 'Période couverte', 'Opérations', 'Lecture', 'Contrôle'],
        [[d['nom'], d['periode'], d['nb_operations'], d['source'],
          {'ok': 'Concordant', 'ecart': f"Écart {d['controle'].get('ecart')} €"}
          .get(d.get('controle', {}).get('statut'), 'Non vérifiable')]
         for d in analyse.get('detail_releves', [])],
        [34, 26, 14, 24, 22])

    fichier = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
    fichier.close()
    try:
        classeur.save(fichier.name)
        with open(fichier.name, 'rb') as f:
            donnees = f.read()
    finally:
        if os.path.exists(fichier.name):
            os.unlink(fichier.name)

    nom = f"depenses_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    reponse = HttpResponse(
        donnees,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    reponse['Content-Disposition'] = f'attachment; filename="{nom}"'
    return reponse


def _nombre_saisi(valeur, defaut):
    """Nombre saisi par l'utilisateur, ou `defaut` si la saisie est illisible ou non finie
    (nan/inf) : ces valeurs feraient planter ou fausseraient les calculs en aval."""
    try:
        v = float(str(valeur).replace(',', '.').replace(' ', ''))
    except (TypeError, ValueError):
        return defaut
    return v if math.isfinite(v) else defaut


def _montant_positif(valeur):
    """Un montant négatif fausserait l'emprunt et la mensualité : ramené à 0."""
    return max(0.0, _nombre_saisi(valeur, 0.0))


def _entier_positif(valeur, defaut):
    """Un effectif négatif n'a pas de sens : on revient à la valeur par défaut."""
    v = _nombre_saisi(valeur, defaut)
    return int(v) if v >= 0 else defaut


def _duree_saisie(valeur, defaut):
    """Une durée hors 5..30 ans fait planter (0) ou fausse (négatif) le calcul de
    mensualité : on revient alors à la durée par défaut."""
    v = _nombre_saisi(valeur, defaut)
    return int(v) if 5 <= v <= 30 else defaut


def _taux_saisi(valeur, defaut):
    """Un taux négatif n'a pas de sens : on revient à la valeur par défaut."""
    v = _nombre_saisi(valeur, defaut)
    return v if v >= 0 else defaut


def _profil_par_defaut():
    return biens_gino.Profil(
        duree=20,
        taux_nominal=TAUX_ACTUELS['regions']['ile_de_france']['20'] + TAUX_ACTUELS['profils']['moyen'],
        taux_assurance=TAUX_ACTUELS['assurance']['30_45'])


def _profil_biens(session):
    """Profil de la page : saisie > simulation > défauts. Renvoie (profil, source, actualise).

    Un taux gardé en session (cookie) depuis un ANCIEN barème est remis au barème
    du jour pour sa durée — règle : toujours les taux les plus récents. `actualise`
    le signale à l'utilisateur ; la session est corrigée pour ne le dire qu'une fois.
    """
    courant = TAUX_DATE.isoformat()
    if session.get('profil_biens'):
        profil = biens_gino.Profil(**session['profil_biens'])
        if session.get('profil_biens_bareme') != courant:
            profil = dataclasses.replace(profil, taux_nominal=taux_pour_duree(profil.duree))
            session['profil_biens'] = {**session['profil_biens'], 'taux_nominal': profil.taux_nominal}
            session['profil_biens_bareme'] = courant
            return profil, 'saisie', True
        return profil, 'saisie', False
    d = _profil_par_defaut()
    if session.get('revenus_nets'):
        actualise = session.get('bareme_date') != courant
        if actualise:
            session['taux_nominal'] = taux_pour_duree(session.get('duree', d.duree))
            session['bareme_date'] = courant
        return biens_gino.Profil(
            revenus=session['revenus_nets'], charges=session.get('charges_fixes', 0),
            apport=session.get('apport', 0), duree=session.get('duree', d.duree),
            taux_nominal=session.get('taux_nominal', d.taux_nominal),
            taux_assurance=session.get('taux_assurance', d.taux_assurance),
            primo=session.get('primo_accedant', False),
            nb_adultes=session.get('nb_adultes', 2), nb_enfants=session.get('nb_enfants', 0)
        ), 'simulation', actualise
    return d, 'defaut', False


def biens_financables(request):
    """Biens relevés par Gino, avec le coût mensuel et le verdict de financement de chacun."""
    racine = Path(settings.AGENCE_IMMO_DIR)
    liste = biens_gino.villes(racine)

    if request.method == 'POST':
        if request.POST.get('action') == 'reprendre':
            request.session.pop('profil_biens', None)
        else:
            d = _profil_par_defaut()
            p = request.POST
            request.session['profil_biens'] = {
                'revenus': _montant_positif(p.get('revenus')),
                'charges': _montant_positif(p.get('charges')),
                'apport': _montant_positif(p.get('apport')),
                'duree': _duree_saisie(p.get('duree'), d.duree),
                'taux_nominal': _taux_saisi(p.get('taux_nominal'), d.taux_nominal),
                'taux_assurance': _taux_saisi(p.get('taux_assurance'), d.taux_assurance),
                'primo': p.get('primo') == 'on',
                'nb_adultes': _entier_positif(p.get('nb_adultes'), 2),
                'nb_enfants': _entier_positif(p.get('nb_enfants'), 0),
            }
            request.session['profil_biens_bareme'] = TAUX_DATE.isoformat()
        query = urlencode({'ville': request.POST.get('ville', '')})
        return redirect(f"{reverse('biens_financables')}?{query}")

    profil, source, taux_actualise = _profil_biens(request.session)
    contexte = {'racine': racine, 'villes': liste or [], 'dossier_introuvable': liste is None,
                'profil': profil, 'source_profil': source, 'taux_actualise': taux_actualise,
                'bareme_durees': bareme_durees()}
    slugs = {v['slug'] for v in liste or []}
    slug = request.GET.get('ville')
    if slug not in slugs:   # ville absente (faute de frappe ou tentative de path traversal) :
        slug = liste[0]['slug'] if liste else None    # on retombe sur la première ville connue
    if slug in slugs:   # jamais de chemin construit à partir d'une valeur non listée
        contexte['ville'] = next(v for v in liste if v['slug'] == slug)
        analyse = biens_gino.analyser_ville(racine / slug, profil, SimulateurPretImmobilier())
        # Affichage seulement (infobulle de l'écart au prix du m²) : aucun calcul de verdict ici.
        aujourd_hui = date.today()
        for b in analyse['biens']:
            b['prix_m2_affiche'] = (round(b['prix'] / b['surface'])
                                    if b.get('prix') and b.get('surface') else None)
            try:
                jours = (aujourd_hui - datetime.fromisoformat(b['verifie_le']).date()).days \
                    if b.get('verifie_le') else None
            except ValueError:
                jours = None
            b['jours_sans_verif'] = jours if jours is not None and jours > 7 else None
            b['pastilles'] = (['neuf' if b.get('neuf') else 'ancien']
                              + (['nouveau'] if b.get('nouveau') else [])
                              + (['baisse'] if b.get('baisse') else []))
            if b.get('baisse'):
                ba, euros = b['baisse'], lambda v: f"{round(v):,}".replace(',', ' ')
                jour = ba['depuis'][8:10] + '/' + ba['depuis'][5:7] if len(ba['depuis']) >= 10 else ''
                pct = f"{abs(ba['pct']):g}".replace('.', ',')
                b['baisse_texte'] = (f"{euros(ba['avant'])} → {euros(ba['apres'])} € (−{pct} %)"
                                     + (f" depuis le {jour}" if jour else ''))
        contexte['analyse'] = analyse
        contexte['mediane_affichee'] = round(analyse['mediane']) if analyse.get('mediane') else None
    # Mensualité que la règle des 35 % laisse disponible (charges déduites), pour l'en-tête.
    contexte['mensualite_max'] = (max(0.0, profil.revenus * biens_gino.SEUIL_FINANCABLE / 100 - profil.charges)
                                  if profil.revenus > 0 else None)
    if slug in slugs:
        dossier = racine / slug
        contexte['veille'] = biens_gino.etat_veille(dossier)
        contexte['veille_en_cours'] = biens_gino.veille_en_cours(dossier)
        contexte['agences_veille'] = biens_gino.agences_veille(dossier)
        fin = (contexte['veille'] or {}).get('fin')
        try:
            contexte['veille_fin'] = datetime.fromisoformat(fin) if fin else None
        except (TypeError, ValueError):
            contexte['veille_fin'] = None
    contexte['veille_impossible'] = request.GET.get('veille') == 'impossible'
    return render(request, 'analyseur/biens_financables.html', contexte)


_pause_lancement = time.sleep


def _demarrer_veille(racine, slug):
    """Lance veille.py d'agence-immo en tâche de fond, détaché de tools_immo.
    tools_immo ne fait que lancer la commande : aucun code d'agence-immo n'est importé."""
    python = Path(settings.AGENCE_IMMO_PYTHON)
    if not python.is_file() or not (racine / 'veille.py').is_file():
        return False
    drapeaux = (getattr(subprocess, 'DETACHED_PROCESS', 0)
                | getattr(subprocess, 'CREATE_NEW_PROCESS_GROUP', 0))
    try:
        with open(racine / slug / '_veille.log', 'w', encoding='utf-8') as journal:
            # « -u » : sortie non bufferisée, le journal se remplit au fil du contrôle.
            subprocess.Popen([str(python), '-u', 'veille.py', '--ville', slug], cwd=str(racine),
                             stdout=journal, stderr=subprocess.STDOUT, creationflags=drapeaux)
    except OSError:
        return False
    # Python met ~1 s à démarrer : on attend le verrou (max 5 s) pour que la page
    # réaffichée après la redirection voie déjà « Vérification en cours ».
    for _ in range(50):
        if biens_gino.veille_en_cours(racine / slug):
            break
        _pause_lancement(0.1)
    return True


def lancer_veille(request):
    """Bouton « Mettre à jour cette ville » : un seul contrôle à la fois par ville."""
    racine = Path(settings.AGENCE_IMMO_DIR)
    slug = request.POST.get('ville', '') if request.method == 'POST' else ''
    slugs = {v['slug'] for v in biens_gino.villes(racine) or []}
    params = {'ville': slug} if slug in slugs else {}
    if slug in slugs and not biens_gino.veille_en_cours(racine / slug):
        if not _demarrer_veille(racine, slug):
            params['veille'] = 'impossible'
    query = urlencode(params)
    return redirect(f"{reverse('biens_financables')}{'?' + query if query else ''}")


def etat_veille_json(request):
    """Avancement du contrôle en cours, lu toutes les 2 s par la page."""
    racine = Path(settings.AGENCE_IMMO_DIR)
    slug = request.GET.get('ville')
    if slug not in {v['slug'] for v in biens_gino.villes(racine) or []}:
        return JsonResponse({}, status=404)
    etat = biens_gino.etat_veille(racine / slug) or {}
    etat['en_cours'] = biens_gino.veille_en_cours(racine / slug)
    return JsonResponse(etat)
